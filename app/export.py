"""Result assembly (shared by the API and exports) and export writers."""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import time

from . import db
from .coding import schema_for_case

VALUE_STATUSES = ("suggested", "derived", "admin_generated", "rule_unclear")
INSUFFICIENT = ("insufficient_evidence", "manual_needed", "not_coded_budget", "model_error", "validation_failed")


def _load(js, default):
    try:
        return json.loads(js) if js else default
    except Exception:
        return default


def review_category(f: dict, sugg: dict | None, review: dict | None) -> str:
    if review and review["action"] in ("accepted", "edited", "cleared"):
        return "confirmed"
    if f.get("rule_missing") or (sugg and sugg["status"] == "rule_missing"):
        return "rule_missing"
    if sugg and (sugg["status"] == "disputed" or (sugg["status"] in VALUE_STATUSES and sugg.get("counter"))):
        return "disputed"
    if sugg and sugg["status"] == "partially_valid":
        return "pending"  # supported selections kept, rejected ones listed; needs a human decision
    if sugg and sugg["status"] in VALUE_STATUSES and sugg["value"]:
        return "pending"
    if review and review["action"] == "deferred":
        return "pending"
    if sugg and sugg["status"] == "analyst_note":
        return "pending"
    return "insufficient"


def case_results(case_id: int) -> dict:
    """Rows for the review table and exports. Each row has the display suggestion (the one provider's result, or a
    neutral combined view when both providers have one — never a preferred provider, D-036), every
    provider's suggestion from the latest coding run ('providers'), the provider comparison, and the human review.
    The human-reviewed value is kept separately from all model suggestions."""
    from .coding import COMPARISON_LABELS, INTERPRETATION_LABELS, PROVIDER_LABELS, interpretation_of, suggestion_sets
    case = db.q1("SELECT * FROM cases WHERE id=?", (case_id,))
    schema = schema_for_case(case)
    sets = suggestion_sets(case_id)
    reviews = {r["variable"]: r for r in db.q("SELECT * FROM reviews WHERE case_id=?", (case_id,))}
    srcs = {s["id"]: s for s in db.q("SELECT id,title,url,final_url,source_type,excluded,relevance_status FROM sources WHERE case_id=?", (case_id,))}
    from .agreement import assess
    from .missingness import FINAL_ACTIONS, consistency_flags
    # D-039: codebook-stated missing-value / consistency rules, checked on human-final values only (never changes them)
    rule_flags = consistency_flags({v: r["value"] for v, r in reviews.items() if r["action"] in FINAL_ACTIONS},
                                   {f["name"]: f for f in schema["fields"]})
    rows = []
    for f in schema["fields"]:
        ss = sets.get(f["name"]) or {}
        s = ss.get("display")
        p = ss.get("previous")
        r = reviews.get(f["name"])
        changed = bool(s and p and (s["value"] != p["value"] or s["status"] != p["status"]))
        comp = ss.get("comparison")
        providers = sorted([{**m, "provider_label": PROVIDER_LABELS.get(m.get("provider"), m.get("provider")),
                             "interpretation": interpretation_of(m),
                             "interpretation_label": INTERPRETATION_LABELS.get(interpretation_of(m), "")}
                            for m in ss.get("members", []) if m.get("provider")], key=lambda m: m["provider"])
        cat = review_category(f, s, r)
        if comp and cat not in ("confirmed", "rule_missing") and comp["status"] not in (
                "model_agreement", "both_insufficient", "evidence_disagreement"):
            cat = "disputed"  # providers disagree (or one output is invalid): route to human review
        if comp and comp["status"] == "evidence_disagreement" and cat not in ("confirmed", "rule_missing"):
            cat = "pending"  # same value, different sources: individual or bulk human review
        comp_out = None
        if comp:
            st = "human_approved" if (r and r["action"] in ("accepted", "edited", "cleared")) else comp["status"]
            comp_out = {**comp, "status": st, "model_status": comp["status"], "label": COMPARISON_LABELS.get(st, st)}
        rows.append({
            "name": f["name"], "position": f["position"], "raw_header": f["raw_header"], "type": f["type"], "multi": f.get("multi"),
            "field_class": f.get("field_class"), "section": f.get("section"), "definition": f.get("definition"),
            "codes": f.get("codes", []), "open_options": f.get("open_options", []), "missing_codes": f.get("missing_codes", []),
            "missing_labels": f.get("missing_labels", {}),
            "codebook_ref": f.get("codebook_ref"), "issues": f.get("issues", []), "rule_missing": f.get("rule_missing"),
            "suggestion": s, "previous": p if changed else None, "review": r, "providers": providers,
            "comparison": comp_out, "category": cat, "cells": _cells_out(ss.get("cells")),
            "agreement": assess(case_id, f, ss, r), "rule_flags": rule_flags.get(f["name"], []),
            "system": ss.get("system") if not providers else None,
        })
    return {"case": case, "schema_label": schema.get("label"), "schema_id": schema.get("id"), "rows": rows, "sources": srcs}


CELL_LABELS = {"not_run": "Not run", "suggested": "Suggested", "disputed": "Disputed", "partially_valid": "Partially valid",
               "no_supported_value": "No supported value", "stopped": "Stopped", "failed": "Failed",
               "invalid_output": "Invalid output", "limit_reached": "Limit reached"}


def _cells_out(cells: dict | None) -> dict:
    """Per-provider cells for the Review table: Claude and OpenAI never share or overwrite a cell."""
    from .coding import interpretation_of as interp_of
    out = {}
    for col in ("anthropic", "openai"):
        c = (cells or {}).get(col) or {"state": "not_run", "row": None}
        row = c.get("row")
        out[col] = {"state": c["state"], "label": CELL_LABELS.get(c["state"], c["state"]),
                    "value": (row or {}).get("value") or "", "status": (row or {}).get("status"),
                    "evidence_status": (row or {}).get("evidence_status"),
                    "evidence_recorded": bool((row or {}).get("evidence_recorded")),
                    "validation_state": (row or {}).get("validation_state"),
                    "suggestion_id": (row or {}).get("id"), "model": (row or {}).get("model") or (row or {}).get("run_model"),
                    "interpretation": interp_of(row) if row else None, "cache_status": (row or {}).get("cache_status"),
                    "evidence_snapshot_id": (row or {}).get("evidence_snapshot_id"),
                    "analysis_spec_id": (row or {}).get("analysis_spec_id"),
                    "stopped_latest": bool(c.get("stopped_latest")),
                    "previous_value": (c.get("previous") or {}).get("value") if c.get("previous") else None,
                    "changed": bool(c.get("previous"))}
    return out


def export_value(row: dict, include_unreviewed: bool) -> tuple[str, str]:
    r = row["review"]
    if r and r["action"] in ("accepted", "edited", "cleared"):
        return r["value"] or "", "reviewed"  # human final always wins (incl. bulk-confirmed agreement)
    s = row["suggestion"]
    comp = row.get("comparison")
    if comp and comp.get("model_status") != "model_agreement":
        return "", "blank"  # providers disagree or an output is invalid: never exported without a human decision
    if include_unreviewed and s and s["status"] in VALUE_STATUSES and s["value"] and not s.get("evidence_review_reason"):
        # D-038: only SUPPORTED (or results predating evidence statuses) are exported unreviewed; INFERRED is not
        return s["value"], "UNREVIEWED"
    return "", "blank"


def source_cell(row: dict, srcs: dict) -> str:
    s = row["suggestion"]
    if not s:
        return ""
    out = []
    for ev in s.get("evidence", []):
        if ev.get("stance") not in ("supports", None):
            continue
        src = srcs.get(ev.get("source_id"), {})
        loc = f"p.{ev['page']}" if ev.get("page") else (f"para {ev['para']}" if ev.get("para") else "")
        out.append(f"{ev['id']} {loc} {src.get('final_url') or src.get('url') or src.get('title') or ''}".strip())
    if not out and s.get("basis") in ("derived", "generated"):
        return f"({s['basis']})"
    return "; ".join(dict.fromkeys(out))


def explanation_cell(row: dict, origin: str) -> str:
    s, r = row["suggestion"], row["review"]
    parts = []
    if origin == "UNREVIEWED":
        parts.append("[UNREVIEWED SUGGESTION — not confirmed by a human coder]")
    if s and s.get("provider") and origin != "reviewed":
        parts.append(f"[model: {s.get('provider')} {s.get('model') or ''}]".replace(" ]", "]"))
    if s and s.get("display_kind") == "two_providers" and origin != "reviewed":
        vb = s.get("values_by_provider") or {}
        parts.append("[independent providers: " + "; ".join(
            f"{'Claude' if p == 'anthropic' else 'OpenAI' if p in ('openai', 'openai_compatible') else p}="
            f"{v or '(blank)'}" for p, v in sorted(vb.items())) + "]")
    if r and r.get("method") == "bulk_independent_agreement":
        parts.append(f"[human-confirmed independent model agreement (bulk confirmation) by {r.get('reviewer') or 'reviewer'}]")
    elif r and r.get("method") == "bulk_separate_run_agreement":  # D-035 (retained): cross-version agreement
        parts.append(f"[human-confirmed CROSS-VERSION model agreement (bulk confirmation, version warning acknowledged) "
                     f"by {r.get('reviewer') or 'reviewer'}]")
    elif r and r.get("source_provider"):
        parts.append(f"[accepted from {r['source_provider']} suggestion]")
    comp = row.get("comparison")
    if comp and comp.get("model_status") != "model_agreement":
        parts.append(f"[providers: {comp.get('model_status')}; human decision needed]")
    if comp and comp.get("kind") == "cross_version":
        parts.append("[cross-version: " + ("same suggested value, different analysis versions"
                                           if comp.get("same_value")
                                           else "results differ, but the models also used different analysis inputs") + "]")
    if r and r["action"] in ("edited", "cleared") and r.get("reason"):
        parts.append(f"Reviewer: {r['reason']}")
    for fl in row.get("rule_flags") or []:
        parts.append(f"[rule check: {fl}]")
    if s:
        if s.get("rationale"):
            parts.append(s["rationale"])
        if s.get("unresolved"):
            parts.append(f"Unresolved: {s['unresolved']}")
        if origin == "blank" and s["status"] not in VALUE_STATUSES:
            parts.append(f"(status: {s['status']})")
        if origin == "blank" and s.get("evidence_review_reason"):
            parts.append(f"({s['evidence_review_reason']})")
    return " ".join(parts)


def evidence_status_text(s: dict | None) -> str:
    """Evidence status for exports; 'not recorded' for results stored before evidence statuses existed (D-038)."""
    if not s or s.get("basis") not in ("model", None) or s.get("status") in ("derived", "admin_generated"):
        return ""
    if s.get("display_kind") == "two_providers":
        return "; ".join(f"{p}: {es or 'not recorded'}" for p, es in sorted((s.get("evidence_statuses") or {}).items()))
    return s.get("evidence_status") or "not recorded"


def _alts_text(s: dict) -> str:
    return "; ".join(f"{a.get('value')} [" + ", ".join(e.get("id", "") for e in a.get("evidence") or []) + "]"
                     for a in s.get("alternatives") or [])


def tsv(case_id: int, include_unreviewed=False) -> str:
    res = case_results(case_id)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter="\t", lineterminator="\n")
    w.writerow(["Variable", "Value", "Source", "Explanation"])
    for row in res["rows"]:
        v, origin = export_value(row, include_unreviewed)
        w.writerow([row["name"], v, source_cell(row, res["sources"]) if v else "", explanation_cell(row, origin)])
    return buf.getvalue()


def _cell_value(v: str, ftype: str):
    if v == "":
        return None
    if ftype == "date" and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        return dt.date.fromisoformat(v)
    if re.fullmatch(r"-?\d+", v) and ftype in ("numeric", "categorical"):
        return int(v)
    if re.fullmatch(r"-?\d+\.\d+", v) and ftype == "numeric":
        return float(v)
    return v  # multi-codes like "1, 3" stay text, so Excel cannot turn them into dates


def xlsx(case_id: int, include_unreviewed=False) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    res = case_results(case_id)
    case = res["case"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    info = [["WFD Coding Assistant export"], ["Case", case["name"]], ["Location", case["location"]], ["Date", case["date_text"]],
            ["Schema version", res["schema_label"]], ["Exported at", dt.datetime.now().isoformat(timespec="seconds")],
            ["Values exported", "Reviewed values + UNREVIEWED model suggestions (marked)" if include_unreviewed else
             "Human-reviewed values only; everything else left blank"],
            ["Note", "Blank cells are genuine blanks (not established or not reviewed). The original workbook was not modified."]]
    for r in info:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)

    wr = wb.create_sheet("Results")
    wr.append(["Variable", "Value", "Source", "Explanation", "Review status", "Value origin", "Suggestion status",
               "Model comparison", "Accepted from provider", "Agreement status", "Confirmation method",
               "Evidence status", "Validation status", "Rule checks"])
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in res["rows"]:
        v, origin = export_value(row, include_unreviewed)
        wr.append([row["name"], v, source_cell(row, res["sources"]) if v else "", explanation_cell(row, origin), row["category"],
                   origin, (row["suggestion"] or {}).get("status", ""), (row.get("comparison") or {}).get("label", ""),
                   (row["review"] or {}).get("source_provider") or "", (row.get("agreement") or {}).get("label", ""),
                   ("human-approved model agreement (bulk)" if (row["review"] or {}).get("method") == "bulk_independent_agreement"
                    else "human-approved cross-version agreement (bulk, D-035)"
                    if (row["review"] or {}).get("method") == "bulk_separate_run_agreement"
                    else ((row["review"] or {}).get("action") or "")),
                   evidence_status_text(row["suggestion"]), (row["suggestion"] or {}).get("validation_state") or "",
                   "; ".join(row.get("rule_flags") or [])])
        if origin == "UNREVIEWED":
            for c in wr[wr.max_row]:
                c.fill = yellow

    wc = wb.create_sheet("Case_Row")
    wc.append([row["raw_header"] for row in res["rows"]])
    wc.append([_cell_value(export_value(row, include_unreviewed)[0], row["type"]) for row in res["rows"]])
    for c in wc[2]:
        if isinstance(c.value, dt.date):
            c.number_format = "yyyy-mm-dd"

    we = wb.create_sheet("Evidence")
    we.append(["Variable", "Stance", "Passage", "Source", "Page", "Paragraph", "Quote", "URL", "Source type", "Provider",
               "Interpretation"])
    for row in res["rows"]:
        members = row.get("providers") or ([row["suggestion"]] if row["suggestion"] else [])
        for s in members:
            for ev in s.get("evidence", []) + s.get("counter", []):
                src = res["sources"].get(ev.get("source_id"), {})
                we.append([row["name"], ev.get("stance"), ev.get("id"), f"S{ev.get('source_id')}", ev.get("page"), ev.get("para"),
                           ev.get("quote"), src.get("final_url") or src.get("url"), src.get("source_type"),
                           s.get("provider") or "", s.get("interpretation") or ""])

    wp = wb.create_sheet("Provider_Suggestions")
    wp.append(["Variable", "Provider", "Model", "Interpretation", "Suggested value", "Status", "Rationale", "Unresolved",
               "Cache status", "Evidence version", "Analysis version", "Run", "Original run", "Generated at",
               "Comparison", "Human-approved value", "Evidence status", "Validation status", "Alternative values",
               "Missing evidence"])
    for row in res["rows"]:
        rv = row["review"] if (row["review"] and row["review"]["action"] in ("accepted", "edited", "cleared")) else None
        for s in row.get("providers") or []:
            ga = s.get("generated_at") or s.get("created_at")
            wp.append([row["name"], s.get("provider"), s.get("model"), s.get("interpretation"), s.get("value"),
                       s.get("status"), s.get("rationale"), s.get("unresolved"), s.get("cache_status") or "",
                       s.get("evidence_snapshot_id") or "not recorded", s.get("analysis_spec_id") or "not recorded",
                       s.get("run_id"), s.get("cache_source_run_id") or s.get("run_id"),
                       dt.datetime.fromtimestamp(ga).isoformat(timespec="seconds") if ga else "",
                       (row.get("comparison") or {}).get("label", ""), (rv or {}).get("value", "") if rv else "",
                       s.get("evidence_status") or "not recorded", s.get("validation_state") or "",
                       _alts_text(s), s.get("missing_evidence") or ""])

    wbk = wb.create_sheet("Bulk_Confirmations")
    bcols = ["batch_id", "variable", "value", "claude_suggestion_id", "openai_suggestion_id", "claude_model", "openai_model",
             "group_id", "codebook_version", "prompt_version", "evidence_difference", "previous_value", "reviewer", "at", "method",
             "evidence_snapshot_id", "analysis_spec_id", "claude_run_id", "openai_run_id", "claude_cache_status",
             "openai_cache_status", "claude_generated_at", "openai_generated_at", "comparison_class", "differences_json",
             "warning_shown", "acknowledged", "claude_versions_json", "openai_versions_json", "selected_json"]
    wbk.append(bcols)
    for bc in db.q("SELECT * FROM bulk_confirmations WHERE case_id=? ORDER BY id", (case_id,)):
        for k in ("at", "claude_generated_at", "openai_generated_at"):
            bc[k] = dt.datetime.fromtimestamp(bc[k]).isoformat(timespec="seconds") if bc.get(k) else ""
        wbk.append([bc.get(k) for k in bcols])

    wsr = wb.create_sheet("Sources")
    cols = ["id", "title", "publisher", "published_date", "source_type", "origin", "found_via", "url", "final_url", "fetch_status",
            "fetch_error", "relevance_status", "relevance_reason", "duplicate_of", "near_duplicate_of", "similarity", "excluded",
            "exclude_reason", "ocr_status", "n_pages", "n_passages", "text_chars", "content_hash", "retrieved_at"]
    wsr.append(cols)
    for s in db.q("SELECT * FROM sources WHERE case_id=? ORDER BY id", (case_id,)):
        s["retrieved_at"] = dt.datetime.fromtimestamp(s["retrieved_at"]).isoformat(timespec="seconds") if s.get("retrieved_at") else ""
        wsr.append([s.get(c) for c in cols])

    wq = wb.create_sheet("Search_Log")
    wq.append(["Round", "Purpose", "Query", "Page", "Provider", "Status", "Results", "Error", "Target variables", "Time"])
    for q in db.q("SELECT * FROM search_queries WHERE case_id=? ORDER BY id", (case_id,)):
        wq.append([q["round"], q["purpose"], q["query"], q["page"], q["provider"], q["status"], q["result_count"], q["error"],
                   q["target_vars"], dt.datetime.fromtimestamp(q["created_at"]).isoformat(timespec="seconds")])

    wu = wb.create_sheet("Runs_Usage")
    wu.append(["Kind", "Provider", "Model", "Input tokens", "Output tokens", "Units", "Cost USD", "Estimated?", "Note", "Time",
               "Reasoning tokens", "Request id"])
    for u in db.q("SELECT * FROM usage WHERE case_id=? ORDER BY id", (case_id,)):
        wu.append([u["kind"], u["provider"], u["model"], u["input_tokens"], u["output_tokens"], u["units"], u["cost_usd"],
                   "yes" if u["estimated"] else "no", u["note"], dt.datetime.fromtimestamp(u["at"]).isoformat(timespec="seconds"),
                   u.get("reasoning_tokens"), u.get("request_id")])
    wu.append([])
    wu.append(["Run id", "Mode", "Model", "Schema version", "Created", "Provider", "Interpretation", "Coding mode", "Group",
               "Evidence version", "Analysis version"])
    from .coding import interpretation_of
    for r in db.q("SELECT * FROM runs WHERE case_id=? ORDER BY id", (case_id,)):
        wu.append([r["id"], r["mode"], r["model"], r["schema_version_id"], dt.datetime.fromtimestamp(r["created_at"]).isoformat(timespec="seconds"),
                   r.get("provider"), interpretation_of(r) if r.get("provider") else "", r.get("coding_mode"),
                   r.get("group_id"), r.get("evidence_snapshot_id") or "", r.get("analysis_spec_id") or ""])

    wm = wb.create_sheet("Field_Mapping")
    wm.append(["Position", "Workbook header", "Codebook variable", "Type", "Multi", "Class", "Codebook reference", "Issues"])
    for row in res["rows"]:
        wm.append([row["position"], row["raw_header"], row["name"] if not row["rule_missing"] else "(no rule)", row["type"],
                   "yes" if row["multi"] else "no", row["field_class"], row["codebook_ref"], " | ".join(row["issues"])])
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
        for col in sh.columns:
            letter = col[0].column_letter
            sh.column_dimensions[letter].width = 18 if sh.title == "Case_Row" else min(70, max(10, max(len(str(c.value or ""))
                                                                                                      for c in col[:50]) + 2))
        if sh.title in ("Results", "Evidence"):
            for r_ in sh.iter_rows(min_row=2):
                for c in r_:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def json_export(case_id: int, include_unreviewed=False) -> dict:
    res = case_results(case_id)
    cited = set()
    for row in res["rows"]:
        s = row["suggestion"]
        if s:
            for ev in s.get("evidence", []) + s.get("counter", []):
                cited.add(ev.get("id"))
    passages = [p for p in db.q("SELECT * FROM passages WHERE case_id=?", (case_id,)) if p["id"] in cited]
    out_rows = []
    for row in res["rows"]:
        v, origin = export_value(row, include_unreviewed)
        out_rows.append({**row, "export_value": v, "export_origin": origin})
    return {
        "export": {"at": dt.datetime.now().isoformat(timespec="seconds"), "include_unreviewed": include_unreviewed,
                   "schema_version": res["schema_label"], "tool": "WFD Coding Assistant"},
        "case": res["case"], "results": out_rows,
        "review_history": db.q("SELECT * FROM review_history WHERE case_id=? ORDER BY id", (case_id,)),
        "bulk_confirmations": db.q("SELECT * FROM bulk_confirmations WHERE case_id=? ORDER BY id", (case_id,)),
        "sources": db.q("SELECT * FROM sources WHERE case_id=? ORDER BY id", (case_id,)),
        "cited_passages": passages,
        "search_log": db.q("SELECT * FROM search_queries WHERE case_id=? ORDER BY id", (case_id,)),
        "usage": db.q("SELECT * FROM usage WHERE case_id=? ORDER BY id", (case_id,)),
        "runs": db.q("SELECT * FROM runs WHERE case_id=? ORDER BY id", (case_id,)),
    }
