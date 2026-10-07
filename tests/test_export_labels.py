"""Readable names in exports and run logs (display only, 2026-10-06). Run:  python -m tests.test_export_labels

Offline: scripted providers (no network), TEST-FIXTURE search. Checks that researchers no longer see raw internal ids
(`anthropic`, `openai`, `dual_independent`, role / comparison / method ids) in the Excel and TSV exports or the run log,
while every internal id, stored value, cache key, the database schema, the JSON export and all coded values stay
exactly the same. Historical `primary` / `reviewer` records keep their recorded role.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import socket
import sys
import tempfile
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="wfdexp_"))
os.environ["WFD_DATA_DIR"] = str(TMP / "data")
for k in ("ANTHROPIC_API_KEY", "TAVILY_API_KEY", "BRAVE_API_KEY", "OPENAI_API_KEY", "OPENAI_BASE_URL", "OPENAI_MODEL",
          "WFD_TEST_FIXTURE"):
    os.environ.pop(k, None)
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
PASS, FAIL, NET = [], [], []
_real = socket.socket.connect


def _guard(self, addr):
    host = addr[0] if isinstance(addr, tuple) else str(addr)
    if host not in ("127.0.0.1", "localhost", "::1"):
        NET.append(host)
        raise OSError("network blocked in tests")
    return _real(self, addr)


socket.socket.connect = _guard
PIN_SCHEMA = "5f231bf5a2af2aad"  # computed on main 302a2f2 / 4924f64 (unchanged since)
PIN_CACHE = {"anthropic": "llm3:ae6ef713619ca7f24bc8bce280acc137c03bf96af88d1493f2a1d3c6b142ad26",
             "openai": "llm3:a82924c4abee492d7d602a6e1ee0445e5a366ca28a047738eda49e597e2c84ab"}
RAW = {"anthropic", "openai", "anthropic+openai", "independent", "cross_model_review", "dual_independent", "anthropic_only",
       "openai_only", "anthropic_primary_openai_review", "single", "tavily", "bulk_independent_agreement",
       "bulk_separate_run_agreement", "matched_version", "cross_version"}
CORPUS = ("Hill Country flash flood, Kerr County, Texas, July 4, 2025. Kerr County officials did not send a CodeRED alert "
          "to residents along the Guadalupe River before the water rose. ")


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  PASS " if cond else "  FAIL ") + name + (f"  — {detail}" if detail and not cond else ""))


class Fake:
    supports_schema = False

    def __init__(s, p, pol):
        s.provider, s.model, s.pol = p, ("claude-sonnet-5-5" if p == "anthropic" else "gpt-6.1-sol"), pol

    def generation_settings(s):
        return {"stand_in": s.provider}

    def check_available(s):
        return None

    def complete(s, system, user, max_tokens=4000):
        names = [l.split("### ")[1].split("  (")[0] for l in user.splitlines() if l.startswith("### ")]
        ps = re.findall(r"^\[(S\d+-P\d+)\] \(source[^\n]*\)\n([^\n]+)", user, re.M)
        res = []
        for n in names:
            v = s.pol.get(n)
            if v is None or not ps:
                res.append({"variable": n, "value": "", "status": "insufficient_evidence", "evidence_status": "INSUFFICIENT",
                            "evidence": []})
            else:
                res.append({"variable": n, "value": v, "status": "suggested", "evidence_status": "SUPPORTED",
                            "rationale": "r", "unresolved": "",
                            "evidence": [{"id": ps[0][0], "quote": ps[0][1].strip()[:60], "stance": "supports"}]})
        return {"text": json.dumps({"results": res}), "input_tokens": 1000, "output_tokens": 200, "stop_reason": "end_turn"}


def main():
    from app import agreement, coding, db, export, research, server
    from app.llm import clients
    from app.llm.structured import build_batch_schema
    pol = {"anthropic": {"SYSTEM_LEVEL": "3", "FAILURE_TYPE": "6", "SYSTEM_INVOLVED": "LOCAL"},
           "openai": {"SYSTEM_LEVEL": "3", "FAILURE_TYPE": "2", "SYSTEM_INVOLVED": "LOCAL"}}
    mk = lambda p, m=None, s=None: Fake(p, pol[p]) if p in pol else None
    clients.make_client = mk
    clients.get_client = lambda p, m, s=None: mk("anthropic" if p in ("auto", "", None) else p)
    os.environ["ANTHROPIC_API_KEY"] = os.environ["OPENAI_API_KEY"] = "test-not-a-real-key"
    fx = TMP / "fx.json"
    fx.write_text(json.dumps({"results": []}))
    os.environ["WFD_TEST_FIXTURE"] = str(fx)
    db.conn()
    server.bootstrap_schema()
    schema = coding.active_schema()
    F = {f["name"]: f for f in schema["fields"]}

    print("\n[1] Internal behaviour fingerprints")
    sql = sorted(r["sql"] or "" for r in db.q("SELECT sql FROM sqlite_master WHERE name NOT LIKE 'sqlite_%'"))
    check("database schema unchanged", hashlib.sha256("\n".join(sql).encode()).hexdigest()[:16] == PIN_SCHEMA)

    class C:
        def __init__(s, p, m):
            s.provider, s.model = p, m

        def generation_settings(s):
            return {"x": 1}
    fs = [F["SYSTEM_LEVEL"], F["SYSTEM_INVOLVED"]]
    pmap = {"S1-P1": {"text": "Kerr County did not send an alert.", "source_id": 1, "page": 1, "para": 1}}
    keys = {p: coding.cache_key(C(p, m), coding.SYSTEM_PROMPT, "PROMPT", fs, ["S1-P1"], pmap,
                                build_batch_schema(fs, ["S1-P1"]) if p == "openai" else None, "1:label")
            for p, m in (("anthropic", "claude-sonnet-5-5"), ("openai", "gpt-6.1-sol"))}
    check("cache keys unchanged (provider ids still 'anthropic' / 'openai' inside the key)", keys == PIN_CACHE)
    check("internal mode names, roles and provider ids unchanged",
          coding.MODES["dual_independent"] == [("anthropic", "independent"), ("openai", "independent")]
          and coding.PUBLIC_MODES == ("anthropic_only", "openai_only", "dual_independent")
          and coding.PROVIDER_LABELS["anthropic"] == "Claude (Anthropic)")

    print("\n[2] A case with a historical cross-model review, a historical 'primary' run and a current dual run")

    class H:
        user = "coder1"
    r = server.api_new_case(H(), {"name": "Hill Country flash flood", "location": "Kerr County, Texas",
                                  "date_text": "July 4, 2025", "start": False, "mode": "dual_independent",
                                  "settings": {"max_search_rounds": 0, "budget_usd": 10, "openai_budget_usd": 10}})
    cid = r["id"]
    research.ingest_manual(cid, "text", {"title": "County statement", "text": CORPUS * 4})
    case = db.q1("SELECT * FROM cases WHERE id=?", (cid,))
    st = json.loads(case["settings_json"])
    coding.ALLOW_CROSS_MODEL_REVIEW = True
    plan, mode, _ = coding.resolve_plan(st, "anthropic_primary_openai_review")
    coding.run_coding_plan(case, st, plan, None, lambda *a: None, None, ["FAILURE_TYPE"], mode=mode)
    coding.ALLOW_CROSS_MODEL_REVIEW = False
    jid = server.api_recode(H(), str(cid), {"variables": ["SYSTEM_LEVEL", "FAILURE_TYPE", "SYSTEM_INVOLVED"],
                                            "mode": "dual_independent"})["job_id"]
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid,)))
    agreement.bulk_confirm(cid, ["SYSTEM_LEVEL"], "coder1")
    server.api_review(H(), str(cid), {"variable": "SYSTEM_INVOLVED", "action": "accept"})
    db.insert("search_queries", {"case_id": cid, "job_id": jid, "round": 0, "purpose": "identify", "query": "q", "page": 1,
                                 "provider": "tavily", "status": "ok", "result_count": 5, "error": None, "target_vars": "",
                                 "created_at": 1791300000.0})
    db.insert("runs", {"case_id": cid, "job_id": None, "mode": "model", "model": "claude-sonnet-5-5",
                       "schema_version_id": case["schema_version_id"], "created_at": 1791200000.0, "provider": "anthropic",
                       "role": "primary", "coding_mode": "single"})
    snap = lambda: json.dumps([db.q(f"SELECT * FROM {t} ORDER BY rowid") for t in
                               ("suggestions", "runs", "reviews", "bulk_confirmations", "model_calls", "usage", "cache")],
                              default=str, sort_keys=True)
    before = snap()

    print("\n[3] Excel export")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(export.xlsx(cid, include_unreviewed=True)))
    after = snap()

    def sheet(name):
        ws = wb[name]
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        return rows[0], rows[1:]
    h, rows = sheet("Provider_Suggestions")
    pcol = [x[h.index("Provider")] for x in rows]
    check("Provider_Suggestions: Provider shows 'Anthropic API' / 'OpenAI API'", set(pcol) == {"Anthropic API", "OpenAI API"},
          str(set(pcol)))
    h, rows = sheet("Bulk_Confirmations")
    check("Bulk_Confirmations: readable, provider-specific headers",
          "Claude model (Anthropic API)" in h and "OpenAI model (OpenAI API)" in h and "Confirmation method" in h
          and not [x for x in h if re.fullmatch(r"[a-z_]+", str(x))], str(h))
    check("Bulk_Confirmations: method and comparison class shown as readable text",
          rows and rows[0][h.index("Confirmation method")] == "Human-approved model agreement (bulk)"
          and rows[0][h.index("Comparison class")] == "Matched analysis version")
    h, rows = sheet("Search_Log")
    check("Search_Log: search service shown as 'Tavily Search API'", "Search service" in h
          and "Tavily Search API" in [x[h.index("Search service")] for x in rows])
    ws = wb["Runs_Usage"]
    all_rows = [[c.value for c in r] for r in ws.iter_rows()]
    i = next(k for k, r in enumerate(all_rows) if r and r[0] == "Run id")
    rh, runs = all_rows[i], [r for r in all_rows[i + 1:] if r and r[0] is not None]
    interp = {r[rh.index("Interpretation")] for r in runs}
    check("Runs_Usage: historical roles kept as recorded ('primary', 'reviewer'), never relabelled as plain independent",
          any(str(x).startswith("primary (historical role)") for x in interp)
          and any(str(x).startswith("reviewer (historical role)") for x in interp), str(interp))
    check("Runs_Usage: coding mode and provider readable",
          {"Claude + OpenAI — independent comparison", "Cross-model review (historical, audit only)",
           "One provider (default of older cases)"} <= {r[rh.index("Coding mode")] for r in runs}
          and {r[rh.index("Provider")] for r in runs} <= {"Anthropic API", "OpenAI API"})
    raw_cells = [(ws_.title, c.coordinate, c.value) for ws_ in wb.worksheets if ws_.title not in ("Field_Mapping",)
                 for row in ws_.iter_rows() for c in row if isinstance(c.value, str) and c.value.strip() in RAW]
    check("no cell anywhere in the Excel export is a raw internal id", not raw_cells, str(raw_cells[:6]))
    h, rows = sheet("Results")
    res = {x[0]: dict(zip(h, x)) for x in rows}
    check("Results: 'Accepted from provider' readable", res["SYSTEM_INVOLVED"]["Accepted from provider"] ==
          "Anthropic API + OpenAI API")
    check("Results: per-provider evidence status names readable",
          res["FAILURE_TYPE"]["Evidence status"] == "Anthropic API: SUPPORTED; OpenAI API: SUPPORTED",
          res["FAILURE_TYPE"]["Evidence status"])
    rr = {r["name"]: r for r in export.case_results(cid)["rows"]}
    check("Results / Case_Row values are the unchanged export values",
          all(str(res[n]["Value"] or "") == str(export.export_value(rr[n], True)[0]) for n in rr)
          and [c.value for c in wb["Case_Row"][2]] == [export._cell_value(export.export_value(rr[n], True)[0], rr[n]["type"])
                                                      for n in [f["name"] for f in schema["fields"]]])

    print("\n[4] TSV, JSON and the run log")
    tsv = list(csv.reader(io.StringIO(export.tsv(cid, include_unreviewed=True)), delimiter="\t"))
    expl = " ".join(r[3] for r in tsv[1:])
    check("TSV explanation: readable provider and comparison names",
          "accepted from Anthropic API + OpenAI API suggestion" in expl and "providers: Value disagreement" in expl
          and "anthropic+openai" not in expl and "value_disagreement" not in expl)
    check("TSV values unchanged", all(r[1] == str(export.export_value(rr[r[0]], True)[0]) for r in tsv[1:]))
    j = export.json_export(cid, include_unreviewed=True)
    check("JSON export keeps the raw internal ids (machine-readable)",
          {x["provider"] for x in j["runs"]} == {"anthropic", "openai"} and
          {x["role"] for x in j["runs"]} == {"independent", "reviewer", "primary"}
          and j["bulk_confirmations"][0]["method"] == "bulk_independent_agreement")
    log = [x["message"] for x in db.q("SELECT message FROM job_log WHERE job_id=?", (jid,))]
    joined = "\n".join(log)
    check("run log: 'Claude + OpenAI — independent comparison' instead of 'mode dual_independent'",
          "Claude + OpenAI — independent comparison: combined worst case" in joined and "mode dual_independent" not in joined)
    check("run log: no raw provider ids in brackets", not re.search(r"\[(anthropic|openai)\]", joined))
    msg = coding.resolve_plan({**st, "model_provider": "auto"}, "nonexistent")[2]
    check("unknown mode error still names the value that was sent", "nonexistent" in (msg or ""))

    print("\n[5] Nothing stored changed")
    check("exporting changed no stored row (suggestions, runs, reviews, bulk confirmations, calls, usage, cache)",
          before == after)
    check("stored provider ids and roles are the raw internal values",
          {(x["provider"], x["role"]) for x in db.q("SELECT provider, role FROM runs")} ==
          {("anthropic", "independent"), ("openai", "reviewer"), ("openai", "independent"), ("anthropic", "primary")})
    check("Human final values unchanged",
          {(x["variable"], x["value"], x["action"]) for x in db.q("SELECT * FROM reviews")} ==
          {("SYSTEM_LEVEL", "3", "accepted"), ("SYSTEM_INVOLVED", "LOCAL", "accepted")})
    check("no network access", not NET, str(NET))
    print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
    if FAIL:
        print("FAILED:", *FAIL, sep="\n  ")
        sys.exit(1)


if __name__ == "__main__":
    main()
