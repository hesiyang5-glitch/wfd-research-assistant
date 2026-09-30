"""Browser-driven end-to-end check (needs Python Playwright + Chromium). Run:  python -m tests.test_ui_flow

Uses the TEST-FIXTURE search provider and a local test web site (no internet). No language model is configured,
so this exercises the manual-coding path through the real UI: form → progress → sources → review → export.
"""
from __future__ import annotations

import glob
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp(prefix="wfdui_"))
SITE = TMP / "site"
SITE.mkdir()
UP = os.environ.get("WFD_TEST_PDFS") or (glob.glob("/root/.claude/uploads/*/") or [""])[0]


def main():
    import http.server
    import threading
    for pat, name in [("DID_WARNINGS", "report.pdf"), ("A_timeline_of_the_catastrophic", "timeline.pdf"), ("Here_s_what_caused", "denver.pdf")]:
        shutil.copy(glob.glob(os.path.join(UP, f"*{pat}*"))[0], SITE / name)

    class H(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **k):
            super().__init__(*a, directory=str(SITE), **k)

        def log_message(self, *a):
            pass
    site = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=site.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{site.server_address[1]}"
    fx = TMP / "fx.json"
    fx.write_text(json.dumps({"results": [
        {"url": f"{base}/report.pdf", "title": "Texas Hill Country floods: did warnings fail victims? 2025", "snippet": "Kerr County July 2025"},
        {"url": f"{base}/timeline.pdf", "title": "Timeline of the Hill Country flash flood, Kerr County Texas 2025", "snippet": "July 4 2025"},
        {"url": f"{base}/denver.pdf", "title": "Hill Country flood 2025 Kerr County alert", "snippet": "Texas 2025"}]}))
    env = {**os.environ, "WFD_DATA_DIR": str(TMP / "data"), "WFD_TEST_FIXTURE": str(fx), "WFD_PORT": "8799"}
    for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "OPENAI_BASE_URL"):
        env.pop(k, None)
    srv = subprocess.Popen([sys.executable, "-m", "app.server"], cwd=str(ROOT), env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    time.sleep(2.5)
    ok = True
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch()
            pg = b.new_page(viewport={"width": 1400, "height": 950}, accept_downloads=True)
            errors = []
            pg.on("pageerror", lambda e: errors.append(str(e)))
            pg.goto("http://127.0.0.1:8799/#/")
            pg.fill("#f_name", "Hill Country flash flood")
            pg.fill("#f_location", "Kerr County, Texas")
            pg.fill("#f_date", "July 4, 2025")
            pg.click("details:nth-of-type(2) summary")
            pg.fill("#s_rounds", "1")
            pg.click("#startBtn")
            pg.wait_for_url("**/progress", timeout=10000)
            for _ in range(120):
                txt = pg.inner_text("#tabBody")
                if "Ready for human review" in txt or "failed" in txt.lower() and "0 failed" not in txt.lower():
                    break
                pg.wait_for_timeout(1000)
            pg.screenshot(path=str(TMP / "ui_progress.png"))
            print("progress done:", "Ready for human review" in txt)
            ok &= "Ready for human review" in txt
            ok &= ("测试夹具" in txt or "TEST FIXTURE" in txt)
            pg.click("text=来源") if "来源" in pg.inner_text(".tabs") else pg.click("text=Sources")
            pg.wait_for_timeout(1200)
            src_txt = pg.inner_text("#tabBody")
            print("irrelevant source flagged:", "irrelevant" in src_txt)
            ok &= "irrelevant" in src_txt
            pg.goto(pg.url.replace("/sources", "/review"))
            pg.wait_for_timeout(1500)
            pg.click("tr[data-v='SYSTEM_LEVEL']")
            pg.wait_for_timeout(500)
            pg.select_option("#edVal", "3")
            pg.fill("#edReason", "County-level alerting authority per S1 p.2")
            pg.click("#editBtn")
            pg.wait_for_timeout(1200)
            row = pg.inner_text("tr[data-v='SYSTEM_LEVEL']")
            print("review saved:", "3" in row)
            ok &= "3" in row
            pg.click("tr[data-v='FAILURE_TYPE']")
            pg.wait_for_timeout(400)
            pg.click("[data-open]")
            pg.wait_for_timeout(800)
            ok &= pg.locator(".passage.hl").count() == 1
            print("citation opens source at highlighted passage:", pg.locator(".passage.hl").count() == 1)
            pg.click("#mClose")
            pg.screenshot(path=str(TMP / "ui_review.png"))
            pg.goto(pg.url.replace("/review", "/export"))
            pg.wait_for_timeout(1000)
            with pg.expect_download() as dl:
                pg.click("text=Excel (XLSX) >> nth=0")
            path = str(TMP / "export.xlsx")
            dl.value.save_as(path)
            import openpyxl
            wb = openpyxl.load_workbook(path)
            cr = list(wb["Case_Row"].iter_rows(values_only=True))
            idx = [c for c in cr[0]].index("SYSTEM_LEVEL")
            print("export has reviewed value:", cr[1][idx] == 3, "| sheets:", wb.sheetnames)
            ok &= cr[1][idx] == 3
            print("JS errors:", errors)
            ok &= not errors
            b.close()
    finally:
        srv.terminate()
        site.shutdown()
    print("UI FLOW:", "PASS" if ok else "FAIL", "| screenshots in", TMP)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
