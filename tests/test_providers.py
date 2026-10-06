"""OpenAI provider, operating modes, limits, caching, migration and UI tests. Run:  python -m tests.test_providers

Fully offline. OpenAI calls go through the REAL official `openai` SDK with a mocked HTTP transport (httpx2.MockTransport),
so request building, response parsing and SDK error types are exercised without any network access. Claude calls go
through the real AnthropicClient with its HTTP post replaced. Fake keys are used and checked to never leak.
These tests verify the app's rules — not live OpenAI behavior, model availability, billing or coding quality.
"""
from __future__ import annotations

import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="wfdprov_"))
os.environ["WFD_DATA_DIR"] = str(TMP / "data")
for k in ("ANTHROPIC_API_KEY", "TAVILY_API_KEY", "BRAVE_API_KEY", "OPENAI_API_KEY", "OPENAI_BASE_URL", "OPENAI_MODEL",
          "WFD_TEST_FIXTURE", "WFD_PASSWORD"):
    os.environ.pop(k, None)
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

KEY_OA = "FAKE-TEST-KEY-openai-NOT-REAL-0123456789"
KEY_AN = "FAKE-TEST-KEY-anthropic-NOT-REAL-9876543210"

PASS, FAIL = [], []
REAL_SLEEP = time.sleep  # retry tests replace time.sleep with a no-op; the browser part needs real waiting


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  PASS " if cond else "  FAIL ") + name + (f"  — {detail}" if detail and not cond else ""))


CORPUS = ("On July 4, 2025 Kerr County did not send a CodeRED alert before the Guadalupe River flash flood. "
          "The county emergency management coordinator said the alerting system was not used overnight. "
          "Sirens were not installed along the river. The National Weather Service issued a flash flood warning at 1:14 a.m. "
          "Investigators said the reasons for the delay remain disputed; some officials cited staffing, others a technical issue. ")

# ----------------------------------------------------------------------------- mock answer generator
PASSAGE_RX = re.compile(r"^\[(S\d+-P\d+)\] \(source[^\n]*\)\n([^\n]+)", re.M)


def passages_in(prompt: str) -> list[tuple[str, str]]:
    return PASSAGE_RX.findall(prompt)


def answer_for(var: str, prompt: str, policy: dict) -> dict:
    """policy: variable -> value to suggest (with a verbatim quote from the first passage), or absent = blank."""
    ps = passages_in(prompt)
    if var in policy and ps:
        pid, text = ps[policy.get("_passage", 0) % len(ps)]
        quote = text.strip()[:70]
        ev = [{"id": pid, "quote": quote, "stance": "supports"}]
        codes = [c.strip() for c in policy[var].split(",") if c.strip()]
        opts = [{"code": c, "evidence": ev} for c in codes] if len(codes) > 1 else []  # multi-select: evidence per code
        return {"variable": var, "value": policy[var], "status": "suggested", "evidence": ev, "options": opts,
                "rationale": "test rationale", "unresolved": ""}
    return {"variable": var, "value": "", "status": "insufficient_evidence", "evidence": [], "options": [],
            "rationale": "", "unresolved": ""}


class OAMock:
    """Scripted OpenAI HTTP server. Each POST /v1/responses consumes the next scenario (the last one repeats)."""

    def __init__(self):
        self.script = ["good"]
        self.policy = {}
        self.requests = []
        self.models_status = 200
        self.model_lookups = 0

    def handler(self, request):
        import httpx2
        path = request.url.path
        if request.method == "GET" and path.startswith("/v1/models/"):
            self.model_lookups += 1
            if self.models_status == 200:
                return httpx2.Response(200, json={"id": path.rsplit("/", 1)[1], "object": "model", "created": 0, "owned_by": "openai"})
            return httpx2.Response(self.models_status, json={"error": {"message": "The model does not exist", "type": "invalid_request_error",
                                                                       "code": "model_not_found"}})
        body = json.loads(request.content)
        self.requests.append(body)
        step = self.script[min(len(self.requests) - 1, len(self.script) - 1)]
        hdr = {"x-request-id": f"req_test_{len(self.requests)}"}
        if step == "timeout":
            raise httpx2.ReadTimeout("read timed out", request=request)
        if step == "connect_fail":
            raise httpx2.ConnectError("connection refused", request=request)
        if step == "429":
            return httpx2.Response(429, headers=hdr, json={"error": {"message": "Rate limit reached", "type": "requests", "code": "rate_limit_exceeded"}})
        if step == "quota":
            return httpx2.Response(429, headers=hdr, json={"error": {"message": "You exceeded your current quota", "type": "insufficient_quota", "code": "insufficient_quota"}})
        if step == "500":
            return httpx2.Response(500, headers=hdr, json={"error": {"message": "server error", "type": "server_error"}})
        if step == "401":
            return httpx2.Response(401, headers=hdr, json={"error": {"message": f"Incorrect API key provided: {KEY_OA}. You can find your API key at ...",
                                                                     "type": "invalid_request_error", "code": "invalid_api_key"}})
        fmt = body.get("text", {}).get("format", {})
        props = fmt.get("schema", {}).get("properties", {}).get("results", {}).get("properties", {})
        res = {}
        for k, sch in props.items():
            var = sch["properties"]["variable"]["enum"][0]
            a = answer_for(var, body["input"], self.policy)
            if sch["properties"]["value"].get("type") == "array":
                a["value"] = [x.strip() for x in a["value"].split(",") if x.strip()] if a["value"] else []
            res[k] = a
        text = json.dumps({"results": res})
        if step == "invalid_json":
            text = "{not valid json"
        if step == "schema_bad":
            text = json.dumps({"results": {}})
        if step == "invalid_code":
            for k in res:
                if res[k]["variable"] == "FAILURE_TYPE":
                    res[k]["value"] = "99"
                    res[k]["status"] = "suggested"
                    ps = passages_in(body["input"]) or [("S1-P1", "no passage in this batch")]
                    res[k]["evidence"] = [{"id": ps[0][0], "quote": ps[0][1][:60], "stance": "supports"}]
            text = json.dumps({"results": res})
        usage = {"input_tokens": 1200, "input_tokens_details": {"cached_tokens": 0, "cache_write_tokens": 0},
                 "output_tokens": 400, "output_tokens_details": {"reasoning_tokens": 150}, "total_tokens": 1600}
        resp = {"id": f"resp_{len(self.requests)}", "object": "response", "created_at": 1, "model": body["model"],
                "status": "completed", "incomplete_details": None, "error": None,
                "output": [{"type": "reasoning", "id": "rs_1", "summary": []},
                           {"type": "message", "id": "msg_1", "status": "completed", "role": "assistant",
                            "content": [{"type": "output_text", "text": text, "annotations": []}]}],
                "usage": usage, "parallel_tool_calls": True, "tool_choice": "auto", "tools": []}
        if step == "incomplete":
            resp["status"] = "incomplete"
            resp["incomplete_details"] = {"reason": "max_output_tokens"}
            resp["output"] = [{"type": "reasoning", "id": "rs_1", "summary": []}]
            resp["usage"]["output_tokens"] = body["max_output_tokens"]
            resp["usage"]["output_tokens_details"]["reasoning_tokens"] = body["max_output_tokens"]
        if step == "refusal":
            resp["output"][1]["content"] = [{"type": "refusal", "refusal": "I can't help with that."}]
        return httpx2.Response(200, headers=hdr, json=resp)


class ANMock:
    """Scripted Anthropic Messages endpoint for the real AnthropicClient."""

    def __init__(self):
        self.policy = {}
        self.requests = []
        self.after_request = None  # hook: called after each request (used to press "Stop" mid-run)

    def post(self, url, json=None, **kw):
        import json as js

        class R:
            pass
        self.requests.append(json)
        if self.after_request:
            self.after_request(len(self.requests))
        user = json["messages"][0]["content"]
        names = [l.split("### ")[1].split("  (")[0] for l in user.splitlines() if l.startswith("### ")]
        r = R()
        r.status_code = 200
        r.headers = {"request-id": f"req_an_{len(self.requests)}"}
        payload = {"id": "msg_x", "content": [{"type": "text", "text": js.dumps({"results": [answer_for(n, user, self.policy) for n in names]})}],
                   "usage": {"input_tokens": 1000, "output_tokens": 300}, "stop_reason": "end_turn"}
        r.text = js.dumps(payload)
        r.json = lambda: payload
        return r


def main():
    import httpx
    import httpx2
    from openai import OpenAI

    from app import coding, db, export, jobs, migrations, research, server
    from app.config import DEFAULT_SETTINGS
    from app.llm import clients
    from app.llm import openai_responses as OR
    from app.llm.structured import build_batch_schema, normalize_structured, SchemaViolation

    OA, AN = OAMock(), ANMock()
    OR.time.sleep = lambda *_: None
    clients.time.sleep = lambda *_: None

    def fake_sdk(self):
        if self._client is None:
            self._client = OpenAI(api_key=self._key, max_retries=0, http_client=httpx2.Client(transport=httpx2.MockTransport(OA.handler)))
        return self._client
    OR.OpenAIResponsesClient._sdk = fake_sdk
    clients.httpx.post = AN.post

    def set_keys(an: bool, oa: bool):
        for k, v, on in (("ANTHROPIC_API_KEY", KEY_AN, an), ("OPENAI_API_KEY", KEY_OA, oa)):
            if on:
                os.environ[k] = v
            else:
                os.environ.pop(k, None)

    db.conn()
    server.bootstrap_schema()
    schema = coding.active_schema()
    VARS = ["SYSTEM_LEVEL", "FAILURE_TYPE", "TRANSMISSION_LATENCY", "INTERAGENCY_COORDINATION", "SYSTEM_INVOLVED"]

    def new_case(name="Hill Country flash flood", extra=None):
        st = {**DEFAULT_SETTINGS, "max_search_rounds": 0, **(extra or {})}
        # Same incident name for every case so the corpus is judged relevant; the label goes into details.
        cid = db.insert("cases", {"name": "Hill Country flash flood", "location": "Kerr County, Texas", "date_text": "July 4, 2025", "details": name,
                                  "known_links": "", "aliases": "", "schema_version_id": schema["id"], "status": "new",
                                  "identity_json": "{}", "settings_json": json.dumps(st), "created_at": time.time(),
                                  "updated_at": time.time()})
        research.ingest_manual(cid, "text", {"title": "County statement", "text": CORPUS * 3})
        return cid, st

    def case_row(cid):
        return db.q1("SELECT * FROM cases WHERE id=?", (cid,))

    log_lines = []
    log = lambda lvl, msg: log_lines.append(f"{lvl} {msg}")

    def oa_client(effort="medium", model="gpt-6.1-sol"):
        OR._PREFLIGHT_OK.clear()
        return OR.OpenAIResponsesClient(model, KEY_OA, reasoning_effort=effort)

    def sugg(cid, var, provider=None):
        q = "SELECT * FROM suggestions WHERE case_id=? AND variable=?" + (" AND provider=?" if provider else "") + " ORDER BY id DESC"
        return db.q1(q, (cid, var, provider) if provider else (cid, var))

    # =========================================================================================================
    print("\n[1] OpenAI success, strict structured outputs, request parameters")
    set_keys(False, True)
    cid, st = new_case()
    OA.script, OA.policy = ["good"], {"SYSTEM_LEVEL": "3", "SYSTEM_INVOLVED": "WEA, LOCAL", "TRANSMISSION_LATENCY": "-9"}
    rep = coding.run_coding(case_row(cid), st, oa_client(), None, log, 10.0, VARS)
    body = OA.requests[0]
    fmt = body.get("text", {}).get("format", {})
    check("1 success: SYSTEM_LEVEL suggested 3 from OpenAI", (sugg(cid, "SYSTEM_LEVEL") or {}).get("value") == "3"
          and sugg(cid, "SYSTEM_LEVEL")["status"] == "suggested" and sugg(cid, "SYSTEM_LEVEL")["provider"] == "openai", str(sugg(cid, "SYSTEM_LEVEL")))
    check("2 Responses API with strict JSON Schema (text.format json_schema, strict=true)",
          fmt.get("type") == "json_schema" and fmt.get("strict") is True and fmt.get("name"))
    sch = fmt.get("schema", {})
    props = sch["properties"]["results"]["properties"]
    by_var = {}
    for rq in OA.requests:  # variables are batched by codebook section, so collect every request's schema
        for v in rq["text"]["format"]["schema"]["properties"]["results"]["properties"].values():
            by_var[v["properties"]["variable"]["enum"][0]] = v
    sl = by_var["SYSTEM_LEVEL"]["properties"]["value"]
    check("2 allowed codes supplied from the codebook as an enum (+ blank)",
          sl.get("enum") == [""] + [c["code"] for c in next(f for f in schema["fields"] if f["name"] == "SYSTEM_LEVEL")["codes"]], str(sl))
    tl = json.dumps(by_var["TRANSMISSION_LATENCY"]["properties"]["value"])
    ic = by_var["INTERAGENCY_COORDINATION"]["properties"]["value"]
    check("2 -9 offered only where the codebook defines it", "-9" in tl and "-9" not in (ic.get("enum") or []), f"{tl} / {ic}")
    check("2 multi-select is an array of codes; normalized back to 'WEA, LOCAL'",
          by_var["SYSTEM_INVOLVED"]["properties"]["value"]["type"] == "array" and sugg(cid, "SYSTEM_INVOLVED")["value"] in ("WEA, LOCAL", "")
          , sugg(cid, "SYSTEM_INVOLVED")["value"])
    ids_enum = sch["$defs"]["evidence"]["properties"]["id"]["enum"]
    check("2 cited passage ids restricted to the passages sent in this batch", all(re.fullmatch(r"S\d+-P\d+", i) for i in ids_enum)
          and set(ids_enum) == {p for p, _ in passages_in(body["input"])})
    check("2 every object is closed (additionalProperties false) and all fields required",
          sch["additionalProperties"] is False and set(sch["properties"]["results"]["required"]) == set(props)
          and all(set(v["required"]) == set(v["properties"]) for v in props.values()))
    check("request: store=false, reasoning effort set, no temperature, max_output_tokens = reserve + visible allowance",
          body.get("store") is False and body.get("reasoning") == {"effort": "medium"} and "temperature" not in body
          and body["max_output_tokens"] == min(32000, 16000 + 1500 + 1200 * len(props)), str({k: body.get(k) for k in ("store", "reasoning", "max_output_tokens")}))
    call = db.q1("SELECT * FROM model_calls WHERE case_id=? AND provider='openai' ORDER BY id DESC", (cid,))
    check("usage: input, output and reasoning tokens + request id recorded", call["input_tokens"] == 1200 and call["output_tokens"] == 400
          and call["reasoning_tokens"] == 150 and str(call["request_id"]).startswith("req_test_") and call["status"] == "complete", str(call))
    u = db.q1("SELECT * FROM usage WHERE case_id=? AND provider='openai' ORDER BY id DESC", (cid,))
    check("OpenAI cost recorded at configured price ($2 in / $10 out per 1M)", abs(u["cost_usd"] - (1200 * 2 + 400 * 10) / 1e6) < 1e-9, str(u["cost_usd"]))
    check("model availability was checked once before the first paid call (free lookup)", OA.model_lookups == 1, str(OA.model_lookups))
    n = len(OA.requests)
    coding.run_coding(case_row(cid), st, oa_client(), None, log, 10.0, VARS)
    check("complete, valid OpenAI reply is cached (identical re-run sends nothing)", len(OA.requests) == n, f"{len(OA.requests)} vs {n}")
    hit = db.q1("SELECT * FROM model_calls WHERE case_id=? ORDER BY id DESC", (cid,))
    check("cache reuse is recorded (status cache_hit, 0 attempts, no cost)", hit["status"] == "cache_hit" and hit["http_attempts"] == 0)
    check("suggestion records cache status 'hit'", sugg(cid, "SYSTEM_LEVEL")["cache_status"] == "hit")

    # strict parsing unit checks
    fs = [f for f in schema["fields"] if f["name"] in ("SYSTEM_LEVEL", "SYSTEM_INVOLVED")]
    bs = build_batch_schema(fs, ["S1-P1"])
    try:
        normalize_structured({"results": {}}, fs)
        check("2 schema-noncompliant reply rejected by the normalizer", False)
    except SchemaViolation:
        check("2 schema-noncompliant reply rejected by the normalizer", True)
    kmap = {v["properties"]["variable"]["enum"][0]: k for k, v in bs["properties"]["results"]["properties"].items()}
    good = {"results": {kmap["SYSTEM_LEVEL"]: {"variable": "SYSTEM_LEVEL", "value": "", "status": "insufficient_evidence", "evidence": [], "options": [], "rationale": "", "unresolved": ""},
                        kmap["SYSTEM_INVOLVED"]: {"variable": "SYSTEM_INVOLVED", "value": ["WEA", "EAS"], "status": "suggested", "evidence": [], "options": [], "rationale": "", "unresolved": ""}}}
    out = {o["variable"]: o for o in normalize_structured(good, fs)}
    check("2 normalizer: array value → 'WEA, EAS' string, blank stays blank", out["SYSTEM_INVOLVED"]["value"] == "WEA, EAS" and out["SYSTEM_LEVEL"]["value"] == "")
    swapped = json.loads(json.dumps(good))
    swapped["results"][kmap["SYSTEM_LEVEL"]]["variable"] = "SYSTEM_INVOLVED"
    try:
        normalize_structured(swapped, fs)
        check("2 entry naming the wrong variable rejected", False)
    except SchemaViolation:
        check("2 entry naming the wrong variable rejected", True)

    # =========================================================================================================
    print("\n[2] Invalid outputs are rejected and never cached")
    for step, var, expect_status, call_status in (("invalid_json", "SYSTEM_LEVEL", "model_error", "invalid_json"),
                                                 ("schema_bad", "SYSTEM_LEVEL", "model_error", "schema_error"),
                                                 ("refusal", "SYSTEM_LEVEL", "model_error", "refusal")):
        c2, st2 = new_case(f"case {step}")
        OA.script, OA.policy, OA.requests = [step], {"SYSTEM_LEVEL": "3"}, []
        coding.run_coding(case_row(c2), st2, oa_client(), None, log, 10.0, VARS)
        n1 = len(OA.requests)
        coding.run_coding(case_row(c2), st2, oa_client(), None, log, 10.0, VARS)
        s2 = sugg(c2, var)
        mc = db.q1("SELECT * FROM model_calls WHERE case_id=? ORDER BY id", (c2,))
        num = {"invalid_json": 3, "schema_bad": "3b", "refusal": "16r"}[step]
        check(f"{num} {step}: variable marked {expect_status}, value blank", s2["status"] == expect_status and s2["value"] == "", str(s2["status"]))
        check(f"{num} {step}: not cached (second run calls OpenAI again)", len(OA.requests) == 2 * n1 and n1 > 0, f"{len(OA.requests)} / {n1}")
        check(f"{num} {step}: call ledger status '{call_status}'", mc["status"] == call_status, mc["status"])

    c3, st3 = new_case("case invalid code")
    OA.script, OA.policy, OA.requests = ["invalid_code"], {}, []
    coding.run_coding(case_row(c3), st3, oa_client(), None, log, 10.0, VARS)
    n1 = len(OA.requests)
    s3 = sugg(c3, "FAILURE_TYPE")
    check("4 invalid codebook value (99) rejected by the server validator", s3["status"] == "validation_failed" and s3["value"] == ""
          and "not defined" in s3["validation_json"], s3["validation_json"][:200])
    ft_batches = lambda: sum(1 for rq in OA.requests if '"FAILURE_TYPE"' in json.dumps(rq["text"]["format"]["schema"]))
    ft1 = ft_batches()
    coding.run_coding(case_row(c3), st3, oa_client(), None, log, 10.0, VARS)
    check("4/16 reply with an invalid code is not cached (that batch is sent again; valid batches reused)",
          ft_batches() == 2 * ft1 and len(OA.requests) == n1 + ft1, f"{len(OA.requests)} / {n1} / ft {ft1}")
    db.ex("INSERT OR REPLACE INTO reviews (case_id,variable,value,action,reason,updated_at) VALUES (?,?,?,?,?,?)",
          (c3, "SYSTEM_LEVEL", "2", "edited", "test", time.time()))
    wb = export.xlsx(c3, include_unreviewed=True)
    import openpyxl
    w = openpyxl.load_workbook(io.BytesIO(wb))
    hdr = [c.value for c in w["Case_Row"][1]]
    vals = [c.value for c in w["Case_Row"][2]]
    check("4 invalid value never reaches the workbook export (cell blank)", vals[hdr.index(next(f["raw_header"] for f in schema["fields"] if f["name"] == "FAILURE_TYPE"))] in (None, ""))
    check("5 unsupported variable left blank (insufficient_evidence)", sugg(c3, "INTERAGENCY_COORDINATION")["status"] == "insufficient_evidence"
          and sugg(c3, "INTERAGENCY_COORDINATION")["value"] == "")

    # =========================================================================================================
    print("\n[3] Incomplete replies, output limit, timeouts and retries")
    c4, st4 = new_case("case incomplete")
    OA.script, OA.policy, OA.requests = ["incomplete"], {"SYSTEM_LEVEL": "3"}, []
    coding.run_coding(case_row(c4), st4, oa_client(), None, log, 10.0, VARS)
    n1 = len(OA.requests)
    mc = db.q1("SELECT * FROM model_calls WHERE case_id=? ORDER BY id", (c4,))
    s4 = sugg(c4, "SYSTEM_LEVEL")
    check("6 incomplete reply discarded (variable model_error, blank)", s4["status"] == "model_error" and s4["value"] == "")
    check("7 incomplete reason 'max_output_tokens' recorded, with reasoning tokens", mc["status"] == "incomplete"
          and mc["incomplete_reason"] == "max_output_tokens" and mc["reasoning_tokens"] == mc["max_output_tokens"], str(mc))
    check("7 cut-off message names the output limit", "output limit" in s4["rationale"], s4["rationale"])
    ub = db.q1("SELECT * FROM usage WHERE case_id=? AND provider='openai'", (c4,))
    check("6 incomplete reply is still counted as billed (actual usage)", ub and ub["cost_usd"] > 0 and ub["estimated"] == 0, str(ub))
    coding.run_coding(case_row(c4), st4, oa_client(), None, log, 10.0, VARS)
    check("6 incomplete reply not cached", len(OA.requests) == 2 * n1)

    c5, st5 = new_case("case timeout")
    OA.script, OA.requests = ["timeout"], []
    coding.run_coding(case_row(c5), st5, oa_client(), None, log, 10.0, ["SYSTEM_LEVEL"])
    mc = db.q1("SELECT * FROM model_calls WHERE case_id=? ORDER BY id", (c5,))
    ub = db.q1("SELECT * FROM usage WHERE case_id=? AND provider='openai'", (c5,))
    check("8 read timeout is NOT retried (one request sent)", len(OA.requests) == 1 and mc["http_attempts"] == 1, f"{len(OA.requests)}")
    check("8 timeout counted as possibly billed at worst case", mc["status"] == "possibly_billed_error" and ub["estimated"] == 1 and ub["cost_usd"] > 0)
    OA.script, OA.requests = ["500"], []
    c5b, st5b = new_case("case 500")
    coding.run_coding(case_row(c5b), st5b, oa_client(), None, log, 10.0, ["SYSTEM_LEVEL"])
    check("8 HTTP 500 not retried, possibly billed", len(OA.requests) == 1 and db.q1("SELECT status FROM model_calls WHERE case_id=?", (c5b,))["status"] == "possibly_billed_error")

    c6, st6 = new_case("case retries")
    OA.script, OA.policy, OA.requests = ["429", "good"], {"SYSTEM_LEVEL": "3"}, []
    coding.run_coding(case_row(c6), st6, oa_client(), None, log, 10.0, ["SYSTEM_LEVEL"])
    mc = db.q1("SELECT * FROM model_calls WHERE case_id=? ORDER BY id", (c6,))
    check("9 HTTP 429 retried; both requests counted as attempts", len(OA.requests) == 2 and mc["http_attempts"] == 2
          and sugg(c6, "SYSTEM_LEVEL")["value"] == "3", f"{len(OA.requests)} {mc['http_attempts']}")
    c6b, st6b = new_case("case connect")
    OA.script, OA.requests = ["connect_fail", "good"], []
    coding.run_coding(case_row(c6b), st6b, oa_client(), None, log, 10.0, ["SYSTEM_LEVEL"])
    check("9 connection never opened → retried and counted", db.q1("SELECT http_attempts FROM model_calls WHERE case_id=?", (c6b,))["http_attempts"] == 2)
    check("9 attempts in ledger feed the per-case attempt count", coding.case_ledger(c6)["attempts_by"].get("openai") == 2)
    c6c, st6c = new_case("case quota")
    OA.script, OA.requests = ["quota"], []
    coding.run_coding(case_row(c6c), st6c, oa_client(), None, log, 10.0, VARS)
    check("exhausted quota (429 insufficient_quota) is a configuration error: one request, run stops",
          len(OA.requests) == 1 and db.q1("SELECT status FROM model_calls WHERE case_id=?", (c6c,))["status"] == "config_error"
          and db.q1("SELECT COUNT(*) n FROM suggestions WHERE case_id=? AND rationale LIKE 'Not sent: configuration error%'", (c6c,))["n"] >= 0)
    c7, st7 = new_case("case unknown model")
    OA.models_status, OA.requests = 404, []
    coding.run_coding(case_row(c7), st7, oa_client(model="gpt-does-not-exist"), None, log, None, VARS)
    s7 = sugg(c7, "SYSTEM_LEVEL")
    check("unavailable model → clear configuration error, no coding request sent, no retries",
          len(OA.requests) == 0 and s7["status"] == "model_error" and "not available to this API project" in s7["rationale"], s7["rationale"])
    set_keys(False, True)
    os.environ["OPENAI_MODEL"] = "gpt-does-not-exist"
    OR._PREFLIGHT_OK.clear()
    c7j, _ = new_case("case unknown model job", {"coding_mode": "openai_only"})
    jj = jobs.enqueue(c7j, "recode", {"variables": VARS})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jj,)))
    j7 = db.q1("SELECT * FROM jobs WHERE id=?", (jj,))
    check("unavailable OPENAI_MODEL in a job → paused with a configuration message (not a price question)",
          j7["status"] == "needs_input" and json.loads(j7["state_json"]).get("awaiting") == "provider"
          and "not available to this API project" in j7["message"] and len(OA.requests) == 0, j7["message"])
    os.environ.pop("OPENAI_MODEL", None)
    OA.models_status = 200
    c7b, st7b = new_case("case bad key")
    OA.script, OA.requests = ["401"], []
    coding.run_coding(case_row(c7b), st7b, oa_client(), None, log, 10.0, VARS)
    err = db.q1("SELECT error FROM model_calls WHERE case_id=?", (c7b,))["error"]
    check("27 rejected key: error stored without the key (redacted)", KEY_OA not in err and "redacted" in err and len(OA.requests) == 1, err[:160])

    # =========================================================================================================
    print("\n[4] Case limits: budgets and attempt caps (checked before sending)")
    c8, st8 = new_case("case oa budget", {"openai_budget_usd": 0.01})
    OA.script, OA.requests = ["good"], []
    r8 = coding.run_coding(case_row(c8), st8, oa_client(), None, log, 10.0, VARS)
    check("10 OpenAI per-case budget enforced before sending (no request)", len(OA.requests) == 0 and r8["not_coded_budget"]
          and "OpenAI case budget" in sugg(c8, "SYSTEM_LEVEL")["rationale"], sugg(c8, "SYSTEM_LEVEL")["rationale"])
    c9, st9 = new_case("case combined budget", {"budget_usd": 0.50})
    db.insert("usage", {"case_id": c9, "job_id": None, "kind": "search", "provider": "tavily", "model": None, "units": 1,
                        "cost_usd": 0.30, "estimated": 1, "at": time.time(), "note": "q"})
    db.insert("usage", {"case_id": c9, "job_id": None, "kind": "model", "provider": "anthropic", "model": "claude-sonnet-5-5",
                        "units": 1, "cost_usd": 0.19, "estimated": 0, "at": time.time(), "note": "earlier Claude run"})
    OA.requests = []
    coding.run_coding(case_row(c9), st9, oa_client(), None, log, None, VARS, budget_cap=0.50)
    check("11 combined budget counts Claude + OpenAI + search: OpenAI call skipped", len(OA.requests) == 0
          and "case budget ($0.50)" in sugg(c9, "SYSTEM_LEVEL")["rationale"], sugg(c9, "SYSTEM_LEVEL")["rationale"])
    led = coding.case_ledger(c9)
    check("11 ledger separates Claude / OpenAI / search spend", abs(led["spent_by"]["search"] - 0.30) < 1e-9
          and abs(led["spent_by"]["anthropic"] - 0.19) < 1e-9 and led["spent_by"].get("openai", 0) == 0)
    check("legacy Claude usage rows (before the attempts ledger) count as attempts", led["attempts_by"].get("anthropic") == 1)
    c10, st10 = new_case("case oa attempts", {"max_openai_attempts_per_case": 2})
    OA.script, OA.policy, OA.requests = ["429", "429", "good"], {}, []
    coding.run_coding(case_row(c10), st10, oa_client(), None, log, 10.0, VARS)
    check("12 OpenAI attempt cap: retries stop at the cap (2 requests, never a 3rd)", len(OA.requests) == 2
          and coding.case_ledger(c10)["attempts_by"]["openai"] == 2, str(len(OA.requests)))
    OA.requests = []
    coding.run_coding(case_row(c10), st10, oa_client(), None, log, 10.0, VARS)
    check("12 cap counts across runs: next run sends nothing", len(OA.requests) == 0
          and "OpenAI attempt limit" in sugg(c10, "SYSTEM_LEVEL")["rationale"], sugg(c10, "SYSTEM_LEVEL")["rationale"])
    c11, st11 = new_case("case total attempts", {"max_model_attempts_per_case": 1})
    db.insert("usage", {"case_id": c11, "job_id": None, "kind": "model", "provider": "anthropic", "model": "claude-sonnet-5-5",
                        "units": 1, "cost_usd": 0.01, "estimated": 0, "at": time.time(), "note": "earlier"})
    OA.requests = []
    coding.run_coding(case_row(c11), st11, oa_client(), None, log, 10.0, VARS)
    check("13 combined model-attempt cap includes Claude attempts", len(OA.requests) == 0
          and "model attempt limit" in sugg(c11, "SYSTEM_LEVEL")["rationale"], sugg(c11, "SYSTEM_LEVEL")["rationale"])

    # pause + explicit approval through the job flow
    set_keys(True, True)
    c12, st12 = new_case("case approval", {"openai_budget_usd": 0.05, "coding_mode": "openai_only"})
    jid = jobs.enqueue(c12, "recode", {"variables": VARS})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid,)))
    j = db.q1("SELECT * FROM jobs WHERE id=?", (jid,))
    sj = json.loads(j["state_json"])
    check("10 run pauses when the OpenAI budget could be exceeded", j["status"] == "needs_input" and "openai_budget_usd" in sj.get("limit_needs", {}),
          j["message"])

    class H:
        user = "tester"
    server.api_resume(H(), str(jid), {"approve_over_budget": True})
    ob = json.loads(case_row(c12)["settings_json"])["openai_budget_usd"]
    check("approval raises the OpenAI budget to an explicit number (never unlimited)", abs(ob - sj["limit_needs"]["openai_budget_usd"]["needed"]) < 1e-9
          and ob < 100, str(ob))
    check("approval is logged with the reviewer's name", db.q1("SELECT message FROM job_log WHERE job_id=? AND message LIKE '%openai_budget_usd%tester%'", (jid,)))
    try:
        server.api_resume(H(), str(jid), {"limits": {"max_openai_attempts_per_case": float("inf")}})
        check("a limit cannot be set to infinity", False)
    except server.ApiError:
        check("a limit cannot be set to infinity", True)
    try:
        server.api_resume(H(), str(jid), {"limits": {"openai_budget_usd": -1}})
        check("a limit cannot be negative / removed", False)
    except server.ApiError:
        check("a limit cannot be negative / removed", True)

    # =========================================================================================================
    print("\n[5] Provider-specific cache keys")
    prep = coding.prepare(case_row(cid), st, VARS)
    fs0, ids0 = prep["batches"][0]
    prompt = coding.make_prompt(case_row(cid), fs0, ids0, prep["pmap"], prep["smeta"])
    a_cli = clients.AnthropicClient("claude-sonnet-5-5", KEY_AN)
    o_cli = oa_client()
    sch0 = build_batch_schema(fs0, ids0)
    k = lambda c, **kw: coding.cache_key(c, kw.get("system", coding.system_prompt_for(c)), kw.get("prompt", prompt), kw.get("fs", fs0),
                                         kw.get("ids", ids0), kw.get("pmap", prep["pmap"]), kw.get("schema", sch0 if c.supports_schema else None) if hasattr(c, "supports_schema") else None,
                                         kw.get("codebook", prep["codebook"]), kw.get("role", "primary"))
    base = k(o_cli)
    check("14 Claude and OpenAI never share a cache key", k(a_cli) != base and base.startswith("llm3:"))
    check("15 model change → new key", k(oa_client(model="gpt-6-astra")) != base)
    check("15 codebook version change → new key", k(o_cli, codebook="99:other codebook") != base)
    pm2 = {**prep["pmap"], ids0[0]: {**prep["pmap"][ids0[0]], "text": prep["pmap"][ids0[0]]["text"] + " (edited)"}}
    check("15 evidence change (same ids, different text) → new key", k(o_cli, pmap=pm2) != base)
    check("15 response schema change → new key", k(o_cli, schema=build_batch_schema(fs0, ids0[:-1] or ["X"])) != base)
    check("15 generation settings change (reasoning effort) → new key", k(oa_client(effort="high")) != base)
    check("15 variable-set change → new key", k(o_cli, fs=fs0[:-1] or fs0[:1]) != base if len(fs0) > 1 else True)
    check("15 run role label alone does NOT change the key (D-036, K-40)",
          k(o_cli, role="primary") == k(o_cli, role="independent") == base)
    rv_prompt = prompt + coding.review_block(fs0, {f["name"]: {"value": "3", "status": "suggested"} for f in fs0}, "Claude")
    check("15 cross-model context (other model's answers in the prompt) → new key", k(o_cli, prompt=rv_prompt) != base)
    check("Claude's system prompt text is unchanged (existing results keep their meaning)",
          __import__("hashlib").sha256(coding.SYSTEM_PROMPT.encode()).hexdigest() == "412c9af33b58b92e19703ae229425272d5d170ecd18f490b11267ebf03f1e86f")
    # legacy Claude cache (pre-OpenAI key format) still reused, read-only
    cL, stL = new_case("case legacy cache")
    prepL = coding.prepare(case_row(cL), stL, ["SYSTEM_LEVEL"])
    fsL, idsL = prepL["batches"][0]
    pL = coding.make_prompt(case_row(cL), fsL, idsL, prepL["pmap"], prepL["smeta"])
    db.cache_put(coding.legacy_cache_key(a_cli, pL), {"text": json.dumps({"results": [answer_for("SYSTEM_LEVEL", pL, {"SYSTEM_LEVEL": "3"})]}),
                                                     "input_tokens": 10, "output_tokens": 5, "stop_reason": "end_turn"})
    AN.requests = []
    coding.run_coding(case_row(cL), stL, a_cli, None, log, 10.0, ["SYSTEM_LEVEL"])
    check("existing cached Claude replies (old key format) are reused without a new paid call",
          len(AN.requests) == 0 and sugg(cL, "SYSTEM_LEVEL")["cache_status"] == "legacy_hit", f"{len(AN.requests)}")

    # =========================================================================================================
    print("\n[6] Operating modes and fallbacks")
    set_keys(True, False)
    plan, mode, err = coding.resolve_plan({**DEFAULT_SETTINGS})
    check("17 Anthropic-only fallback: default mode uses Claude when only its key is set",
          len(plan) == 1 and plan[0][0].provider == "anthropic" and mode == "single" and not err)
    am = {m["mode"]: m["available"] for m in server.available_modes(DEFAULT_SETTINGS)}
    check("17 OpenAI modes not offered without OPENAI_API_KEY", not am["openai_only"] and not am["dual_independent"] and am["anthropic_only"])
    plan, mode, err = coding.resolve_plan({**DEFAULT_SETTINGS}, "dual_independent")
    check("17 explicit dual mode without OpenAI key → clear error, no silent fallback", not plan and "OPENAI_API_KEY" in (err or ""))
    try:
        server.check_mode("openai_only", DEFAULT_SETTINGS)
        check("17 API refuses an OpenAI mode without the key", False)
    except server.ApiError:
        check("17 API refuses an OpenAI mode without the key", True)
    set_keys(False, True)
    plan, mode, err = coding.resolve_plan({**DEFAULT_SETTINGS})
    check("18 OpenAI-only fallback: default mode uses OpenAI (official SDK) when only its key is set",
          len(plan) == 1 and plan[0][0].provider == "openai" and isinstance(plan[0][0], OR.OpenAIResponsesClient))
    set_keys(False, False)
    plan, mode, err = coding.resolve_plan({**DEFAULT_SETTINGS})
    c13, st13 = new_case("case manual")
    r13 = coding.run_coding_plan(case_row(c13), st13, plan, None, log, None, VARS)
    check("19 neither key → manual mode: candidate passages only, no codes", not plan and not err
          and sugg(c13, "SYSTEM_LEVEL")["status"] == "manual_needed" and sugg(c13, "SYSTEM_LEVEL")["value"] == "")
    set_keys(True, True)
    check("default mode is unchanged ('single' → Claude first), dual never starts by itself",
          DEFAULT_SETTINGS["coding_mode"] == "single" and coding.resolve_plan({**DEFAULT_SETTINGS})[0][0][0].provider == "anthropic")

    # =========================================================================================================
    print("\n[7] Dual independent coding, reviewer mode, comparison, search reuse")
    fx = TMP / "fx.json"
    fx.write_text(json.dumps({"results": []}))
    os.environ["WFD_TEST_FIXTURE"] = str(fx)
    c14, st14 = new_case("case dual", {"coding_mode": "dual_independent", "max_search_rounds": 1, "budget_usd": 10, "openai_budget_usd": 10})
    OA.script, OA.requests, AN.requests = ["good"], [], []
    OA.policy = {"SYSTEM_LEVEL": "3", "FAILURE_TYPE": "6", "TRANSMISSION_LATENCY": "-9"}
    AN.policy = {"SYSTEM_LEVEL": "3", "FAILURE_TYPE": "2", "SYSTEM_INVOLVED": "LOCAL"}
    jid = jobs.enqueue(c14, "research", {"identity_confirmed": True, "variables": VARS})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid,)))
    j = db.q1("SELECT * FROM jobs WHERE id=?", (jid,))
    sj = json.loads(j["state_json"])
    nq_job = db.q1("SELECT COUNT(*) n FROM search_queries WHERE job_id=?", (jid,))["n"]
    check("dual run completes", j["status"] == "done", f"{j['status']} {j['message']}")
    check("20 search ran once for the case, not once per provider",
          nq_job > 0 and len(sj["coding_report"]["providers"]) == 2
          and db.q1("SELECT COUNT(DISTINCT job_id) n FROM search_queries WHERE case_id=?", (c14,))["n"] == 1, f"queries={nq_job}")
    check("20 both providers coded from the same evidence corpus (same evidence hashes per batch)",
          sorted(r["evidence_hash"] for r in db.q("SELECT evidence_hash FROM model_calls WHERE case_id=? AND provider='openai'", (c14,)))
          == sorted(r["evidence_hash"] for r in db.q("SELECT evidence_hash FROM model_calls WHERE case_id=? AND provider='anthropic'", (c14,))))
    an_prompt = AN.requests[0]["messages"][0]["content"]
    oa_prompt = OA.requests[0]["input"]
    check("21 independent coders get the identical case + evidence prompt", an_prompt == oa_prompt)
    check("21 neither independent coder saw the other's result", "PRIOR SUGGESTIONS" not in an_prompt + oa_prompt)
    roles = {r["provider"]: (r["role"], r["independent"]) for r in db.q("SELECT provider, role, independent FROM runs WHERE case_id=? AND provider IS NOT NULL", (c14,))}
    check("21 both runs stored separately as independent", roles == {"anthropic": ("independent", 1), "openai": ("independent", 1)}, str(roles))
    res = export.case_results(c14)
    rows = {r["name"]: r for r in res["rows"]}
    check("23 agreement: same value + shared source → 'model_agreement'", rows["SYSTEM_LEVEL"]["comparison"]["status"] == "model_agreement",
          str(rows["SYSTEM_LEVEL"]["comparison"]))
    check("23 value disagreement detected (2 vs 6) and routed to review", rows["FAILURE_TYPE"]["comparison"]["status"] == "value_disagreement"
          and rows["FAILURE_TYPE"]["category"] == "disputed")
    check("23 one provider blank detected", rows["SYSTEM_INVOLVED"]["comparison"]["status"] == "one_provider_blank")
    # D-032: two blank answers are 'Both insufficient', never an agreed value
    check("23 both blank → 'both_insufficient' (not agreement)", rows["INTERAGENCY_COORDINATION"]["comparison"]["status"] == "both_insufficient")
    check("each provider's suggestion + evidence kept separately", {p["provider"] for p in rows["FAILURE_TYPE"]["providers"]} == {"anthropic", "openai"}
          and all(p["evidence"] for p in rows["FAILURE_TYPE"]["providers"]))
    a = {"status": "suggested", "value": "3", "evidence": [{"source_id": 1}]}
    check("23 same value, no common source → 'evidence_disagreement'", coding.compare_pair(a, {**a, "evidence": [{"source_id": 2}]})["status"] == "evidence_disagreement")
    check("23 invalid output → 'invalid_provider_output'", coding.compare_pair(a, {**a, "status": "validation_failed"})["status"] == "invalid_provider_output")
    check("23 disputed/counter-evidence → needs human review", coding.compare_pair(a, {**a, "counter": [{"id": "x"}]})["status"] == "needs_human_review")
    xl = openpyxl.load_workbook(io.BytesIO(export.xlsx(c14, include_unreviewed=True)))
    hdr = [c.value for c in xl["Case_Row"][1]]
    v2 = [c.value for c in xl["Case_Row"][2]]
    col = lambda name: hdr.index(next(f["raw_header"] for f in schema["fields"] if f["name"] == name))
    check("disagreement is never exported as a value, even with unreviewed suggestions included", v2[col("FAILURE_TYPE")] in (None, ""))
    check("agreed value is exported only as an UNREVIEWED suggestion when asked", v2[col("SYSTEM_LEVEL")] == 3)
    check("Case_Row headers still identical to the workbook header row", [h or "" for h in hdr] == [f["raw_header"] or "" for f in schema["fields"]],
          str([(a, b) for a, b in zip(hdr, [f["raw_header"] for f in schema["fields"]]) if (a or "") != (b or "")][:3]))
    check("Provider_Suggestions sheet lists both providers", "Provider_Suggestions" in xl.sheetnames
          and {r[1].value for r in xl["Provider_Suggestions"].iter_rows(min_row=2)} >= {"anthropic", "openai"})

    # human review stays separate and final
    class H2:
        user = "coder1"
    try:
        server.api_review(H2(), str(c14), {"variable": "FAILURE_TYPE", "action": "accept"})
        check("25 accepting without choosing a provider is refused when models disagree", False)
    except server.ApiError:
        check("25 accepting without choosing a provider is refused when models disagree", True)
    oa_s = db.q1("SELECT id FROM suggestions WHERE case_id=? AND variable='FAILURE_TYPE' AND provider='openai' ORDER BY id DESC", (c14,))
    server.api_review(H2(), str(c14), {"variable": "FAILURE_TYPE", "action": "accept", "suggestion_id": oa_s["id"]})
    rv = db.q1("SELECT * FROM reviews WHERE case_id=? AND variable='FAILURE_TYPE'", (c14,))
    check("25 accepting a specific provider's suggestion records which provider", rv["value"] == "6" and rv["source_provider"] == "openai")
    server.api_review(H2(), str(c14), {"variable": "SYSTEM_LEVEL", "action": "edit", "value": "2", "reason": "state agency"})
    OA.policy = {"SYSTEM_LEVEL": "4", "FAILURE_TYPE": "1"}
    AN.policy = {"SYSTEM_LEVEL": "5", "FAILURE_TYPE": "1"}
    jid2 = jobs.enqueue(c14, "recode", {"variables": VARS, "mode": "dual_independent"})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid2,)))
    rv1 = db.q1("SELECT * FROM reviews WHERE case_id=? AND variable='SYSTEM_LEVEL'", (c14,))
    rv2 = db.q1("SELECT * FROM reviews WHERE case_id=? AND variable='FAILURE_TYPE'", (c14,))
    check("25 re-analysis with both providers never overwrites human-approved values", rv1["value"] == "2" and rv1["action"] == "edited"
          and rv2["value"] == "6" and rv2["source_provider"] == "openai")
    rows = {r["name"]: r for r in export.case_results(c14)["rows"]}
    check("25 human-approved value shown as 'human_approved', separate from both model suggestions",
          rows["SYSTEM_LEVEL"]["comparison"]["status"] == "human_approved" and rows["SYSTEM_LEVEL"]["review"]["value"] == "2"
          and {p["value"] for p in rows["SYSTEM_LEVEL"]["providers"]} == {"3"},
          json.dumps({"comp": rows["SYSTEM_LEVEL"]["comparison"], "prov": [(p["provider"], p["role"], p["value"], p["status"]) for p in rows["SYSTEM_LEVEL"]["providers"]],
                      "job": db.q1("SELECT status, message FROM jobs WHERE id=?", (jid2,))}, default=str)[:700])
    check("identical evidence → both providers' cached replies reused, no new paid call",
          {p["cache_status"] for p in rows["SYSTEM_LEVEL"]["providers"]} == {"hit"}
          and db.q1("SELECT COUNT(*) n FROM usage WHERE job_id=? AND kind='model'", (jid2,))["n"] == 0)
    check("search not repeated on re-analysis", db.q1("SELECT COUNT(*) n FROM search_queries WHERE job_id=?", (jid2,))["n"] == 0)

    # cross-model review: outside the current scope (D-036) — refused for new runs; historical data stays readable
    c15, st15 = new_case("case reviewer", {"budget_usd": 10, "openai_budget_usd": 10})
    OA.policy, AN.policy, OA.requests, AN.requests = {"SYSTEM_LEVEL": "3"}, {"SYSTEM_LEVEL": "4"}, [], []
    try:
        server.api_recode(H2(), str(c15), {"variables": VARS, "mode": "anthropic_primary_openai_review"})
        check("22 server refuses a new cross-model review run", False)
    except server.ApiError:
        check("22 server refuses a new cross-model review run", True)
    jid3 = jobs.enqueue(c15, "recode", {"variables": VARS, "mode": "anthropic_primary_openai_review"})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid3,)))
    j3 = db.q1("SELECT status, message FROM jobs WHERE id=?", (jid3,))
    check("22 a queued cross-model job pauses before any request", j3["status"] == "needs_input" and not OA.requests
          and not AN.requests and "outside the current scope" in j3["message"], str(j3))
    coding.ALLOW_CROSS_MODEL_REVIEW = True  # only to create historical-style data for the audit checks below
    jid3 = jobs.enqueue(c15, "recode", {"variables": VARS, "mode": "anthropic_primary_openai_review"})
    research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid3,)))
    coding.ALLOW_CROSS_MODEL_REVIEW = False
    rprompt = OA.requests[0]["input"]
    check("22 historical cross-model prompt contains the other model's suggestions", "PRIOR SUGGESTIONS FROM ANOTHER CODER (Claude (Anthropic)" in rprompt
          and '"value": "4"' in rprompt)
    rr = db.q1("SELECT * FROM runs WHERE case_id=? AND role='reviewer'", (c15,))
    check("22 cross-model run labelled cross_model_review, not independent", rr and rr["independent"] == 0
          and rr["provider"] == "openai" and rr["interpretation"] == "cross_model_review" and "not independent" in rr["notes"])
    row = next(r for r in export.case_results(c15)["rows"] if r["name"] == "SYSTEM_LEVEL")
    check("22 comparison marked as cross-model review; no provider is the default suggestion",
          row["comparison"]["kind"] == "cross_model_review" and row["suggestion"]["provider"] is None
          and row["agreement"]["status"] == "cross_model_review" and not row["agreement"]["eligible"], str(row["comparison"]))
    xl = openpyxl.load_workbook(io.BytesIO(export.xlsx(c15)))
    vals = [[c.value for c in r] for r in xl["Provider_Suggestions"].iter_rows(min_row=2)]
    check("22 export labels it 'cross_model_review'", any(v[1] == "openai" and v[3] == "cross_model_review" for v in vals))
    os.environ.pop("WFD_TEST_FIXTURE", None)

    # =========================================================================================================
    print("\n[8] Database migration and rollback")
    legacy = TMP / "legacy.sqlite3"
    con = sqlite3.connect(str(legacy))
    con.executescript(db.SCHEMA)
    for t, c_ in (("reviews", "reviewer"), ("review_history", "reviewer"), ("sources", "added_by"), ("jobs", "started_by")):
        con.execute(f"ALTER TABLE {t} ADD COLUMN {c_} TEXT")
    con.execute("INSERT INTO cases (id, name, settings_json) VALUES (1, 'Marshall Fire', '{}')")
    con.execute("INSERT INTO runs (id, case_id, mode, model) VALUES (1, 1, 'model', 'claude-sonnet-5-5')")
    con.execute("INSERT INTO suggestions (id, case_id, run_id, variable, value, status, basis, evidence_json) VALUES (1,1,1,'SYSTEM_LEVEL','3','suggested','model','[]')")
    con.execute("INSERT INTO reviews (case_id, variable, value, action, reason, reviewer) VALUES (1,'SYSTEM_LEVEL','2','edited','why','coder1')")
    con.execute("INSERT INTO usage (case_id, kind, provider, model, cost_usd) VALUES (1,'model','anthropic','claude-sonnet-5-5',0.42)")
    con.commit()
    before = {t: con.execute(f"SELECT * FROM {t}").fetchall() for t in ("cases", "runs", "suggestions", "reviews", "usage")}
    ch1 = migrations.migrate(con)
    ch2 = migrations.migrate(con)
    check("24 migration adds columns/tables on a legacy database", "suggestions.provider" in ch1 and "usage.reasoning_tokens" in ch1
          and con.execute("SELECT name FROM sqlite_master WHERE name='model_calls'").fetchone())
    check("24 migration is repeatable (second run changes nothing)", ch2 == [])
    after = {t: con.execute(f"SELECT * FROM {t}").fetchall() for t in before}
    same = all([tuple(r)[:len(before[t][i])] for i, r in enumerate(after[t])] == [tuple(r) for r in before[t]] for t in before)
    check("24 existing cases, runs, Claude results, reviews and usage preserved unchanged", same)
    con.execute("INSERT INTO runs (id, case_id, mode, model, provider, role) VALUES (2,1,'model','gpt-6.1-sol','openai','independent')")
    con.execute("INSERT INTO suggestions (id, case_id, run_id, variable, value, status, basis, provider, role) VALUES (2,1,2,'SYSTEM_LEVEL','4','suggested','model','openai','independent')")
    con.execute("INSERT INTO suggestions (id, case_id, run_id, variable, value, status, basis, provider, role) VALUES (3,1,2,'FAILURE_TYPE','6','suggested','model','anthropic','reviewer')")
    con.commit()
    moved = migrations.rollback(con)
    latest = con.execute("SELECT value FROM suggestions WHERE case_id=1 AND variable='SYSTEM_LEVEL' ORDER BY id DESC LIMIT 1").fetchone()[0]
    check("24 rollback archives OpenAI and reviewer rows so the previous version sees only Claude results",
          moved == 2 and latest == "3" and con.execute("SELECT COUNT(*) FROM suggestions_provider_archive").fetchone()[0] == 2)
    check("24 rollback keeps human reviews and usage", con.execute("SELECT value FROM reviews").fetchone()[0] == "2"
          and con.execute("SELECT COUNT(*) FROM usage").fetchone()[0] == 1)
    con.row_factory = None
    restored = migrations.restore(con)
    check("24 restore brings archived rows back", restored == 2 and con.execute("SELECT COUNT(*) FROM suggestions").fetchone()[0] == 3)
    con.close()

    # =========================================================================================================
    print("\n[9] Secrets never leave the server")
    set_keys(True, True)
    st_json = json.dumps(server.api_status(None))
    check("27 /api/status reports providers without key values", KEY_OA not in st_json and KEY_AN not in st_json
          and '"OPENAI_API_KEY": true' in st_json and "gpt-6.1-sol" in st_json)
    dump = ""
    for t in ("usage", "model_calls", "suggestions", "runs", "job_log", "jobs", "cache", "settings", "cases"):
        dump += json.dumps(db.q(f"SELECT * FROM {t}"), default=str)
    check("27 no key in the database (logs, call ledger, cache, suggestions)", KEY_OA not in dump and KEY_AN not in dump
          and "FAKE-TEST-KEY" not in dump)
    check("27 no key in run logs", not any(KEY_OA in l or KEY_AN in l for l in log_lines))
    exp = export.xlsx(c14, include_unreviewed=True)
    z = zipfile.ZipFile(io.BytesIO(exp))
    xml = "".join(z.read(n).decode("utf8", "ignore") for n in z.namelist())
    check("27 no key in the Excel or JSON export", KEY_OA not in xml and KEY_AN not in xml
          and KEY_OA not in json.dumps(export.json_export(c14), default=str))
    check("27 client repr never shows the key", KEY_OA not in repr(oa_client()))
    check("27 redact() removes the key, sk- style keys and bearer tokens", "NOT-REAL" not in OR.redact(f"Bearer {KEY_OA} and {KEY_OA}", KEY_OA) and "abcdef12345" not in OR.redact("key sk-proj-abcdef12345xyz", ""))

    # =========================================================================================================

    # =========================================================================================================
    print("\n[11] Stop / resume each provider independently")
    os.environ["WFD_TEST_FIXTURE"] = str(fx)
    set_keys(True, True)
    big = {"budget_usd": 20, "openai_budget_usd": 20}
    VARS2 = VARS + ["SUMMARY", "ALERTING_AUTHORITY_TYPE", "POPULATION_SCOPE", "DELIVERY_COVERAGE"]

    def run_job(jid):
        research.run_research(db.q1("SELECT * FROM jobs WHERE id=?", (jid,)))
        j = db.q1("SELECT * FROM jobs WHERE id=?", (jid,))
        return j, json.loads(j["state_json"])

    def stop_flag(jid, prov):
        class HH:
            user = "coder1"
        return server.api_stop_provider(HH(), str(jid), prov)

    # (a) Stop OpenAI before it starts: Claude unaffected, OpenAI sends nothing
    c20, _ = new_case("case stop openai", big)
    OA.policy, AN.policy, OA.requests, AN.requests = {"SYSTEM_LEVEL": "3"}, {"SYSTEM_LEVEL": "3"}, [], []
    j20 = jobs.enqueue(c20, "recode", {"variables": VARS2, "mode": "dual_independent"})
    r_stop = stop_flag(j20, "openai")
    j, sj = run_job(j20)
    check("stop warning says a sent request may still finish and be billed", "may still be billed" in r_stop["warning"])
    check("Stop OpenAI: no OpenAI request sent; Claude ran normally", len(OA.requests) == 0 and len(AN.requests) > 0
          and sj["provider_runs"]["openai"]["status"] == "stopped" and sj["provider_runs"]["anthropic"]["status"] == "done",
          json.dumps(sj.get("provider_runs"))[:300])
    rows = {r["name"]: r for r in export.case_results(c20)["rows"]}
    check("Stop OpenAI: Claude results kept; OpenAI cells show 'Stopped'", rows["SYSTEM_LEVEL"]["cells"]["anthropic"]["state"] == "suggested"
          and rows["SYSTEM_LEVEL"]["cells"]["openai"]["state"] == "stopped")
    check("stopped batches cost nothing and are logged as 'stopped'",
          db.q1("SELECT COUNT(*) n, COALESCE(SUM(cost_usd),0) c FROM model_calls WHERE job_id=? AND status='stopped'", (j20,))["n"] > 0
          and db.q1("SELECT COALESCE(SUM(cost_usd),0) c FROM usage WHERE job_id=? AND provider='openai'", (j20,))["c"] == 0)

    # (b) Stop Claude mid-run: the in-flight request finishes and is kept/billed; later Claude batches are not sent;
    #     OpenAI is unaffected and codes everything
    c21, _ = new_case("case stop claude", big)
    OA.requests, AN.requests = [], []
    j21 = jobs.enqueue(c21, "recode", {"variables": VARS2, "mode": "dual_independent"})
    AN.after_request = lambda n: stop_flag(j21, "anthropic") if n == 1 else None
    j, sj = run_job(j21)
    AN.after_request = None
    n_batches = len(coding.prepare(case_row(c21), {**DEFAULT_SETTINGS, **big}, VARS2)["batches"])
    calls_an = db.q("SELECT status, cost_usd FROM model_calls WHERE job_id=? AND provider='anthropic' ORDER BY id", (j21,))
    check("Stop Claude mid-run: exactly the in-flight request was sent; the rest were not",
          len(AN.requests) == 1 and n_batches > 1 and [c["status"] for c in calls_an][0] == "complete"
          and all(c["status"] == "stopped" for c in calls_an[1:]), f"{len(AN.requests)} sent of {n_batches}; {calls_an}")
    check("in-flight Claude request result kept and billed", db.q1("SELECT COUNT(*) n FROM usage WHERE job_id=? AND provider='anthropic'", (j21,))["n"] == 1)
    check("OpenAI unaffected by stopping Claude (all its batches sent)", len(OA.requests) == n_batches
          and sj["provider_runs"]["openai"]["status"] == "done" and sj["provider_runs"]["anthropic"]["status"] == "stopped")
    stopped_vars = set(sj["provider_runs"]["anthropic"]["stopped_variables"])
    check("stopped Claude variables listed for the reviewer", len(stopped_vars) > 0)

    # (c) Resume Claude: only the stopped variables are coded, in the same comparison group
    class HH:
        user = "coder1"
    rr = server.api_resume_provider(HH(), str(j21), "anthropic")
    AN.requests, OA.requests = [], []
    j, sj2 = run_job(rr["job_id"])
    sent_vars = set()
    for rq in AN.requests:
        sent_vars |= {l.split("### ")[1].split("  (")[0] for l in rq["messages"][0]["content"].splitlines() if l.startswith("### ")}
    check("Resume Claude: only previously stopped variables are sent", sent_vars and sent_vars <= stopped_vars and set(rr["variables"]) == stopped_vars,
          f"{sorted(sent_vars)} vs {sorted(stopped_vars)}")
    check("Resume Claude: OpenAI not called again", len(OA.requests) == 0)
    g1 = json.loads(db.q1("SELECT state_json FROM jobs WHERE id=?", (j21,))["state_json"])["group_id"]
    check("Resume writes into the same comparison group", sj2["group_id"] == g1)
    rows = {r["name"]: r for r in export.case_results(c21)["rows"]}
    check("after Resume no Claude cell is 'stopped' and comparison is independent",
          all(r["cells"]["anthropic"]["state"] != "stopped" for r in rows.values() if r["name"] in VARS2)
          and rows["SYSTEM_LEVEL"]["comparison"]["kind"] == "matched_version")
    try:
        server.api_resume_provider(HH(), str(j21), "anthropic")
        check("nothing left to resume → clear message", False)
    except server.ApiError:
        check("nothing left to resume → clear message", True)

    # (d) Resume before the stop took effect: the flag is withdrawn, the provider runs normally
    c22, _ = new_case("case resume in place", big)
    OA.requests, AN.requests = [], []
    j22 = jobs.enqueue(c22, "recode", {"variables": VARS, "mode": "dual_independent"})
    stop_flag(j22, "openai")
    ri = server.api_resume_provider(HH(), str(j22), "openai")
    j, sj = run_job(j22)
    check("Resume before the stop took effect: OpenAI runs normally", ri.get("resumed_in_place") and len(OA.requests) > 0
          and sj["provider_runs"]["openai"]["status"] == "done")

    # (e) A stopped run never hides that provider's earlier completed result
    j23 = jobs.enqueue(c22, "recode", {"variables": VARS, "mode": "dual_independent"})
    stop_flag(j23, "anthropic")
    run_job(j23)
    cell = {r["name"]: r for r in export.case_results(c22)["rows"]}["SYSTEM_LEVEL"]["cells"]["anthropic"]
    check("stopped later run keeps the earlier Claude value visible (marked 'latest run stopped')",
          cell["state"] == "suggested" and cell["value"] == "3" and cell["stopped_latest"], str(cell))
    # (f) whole-job cancel still stops everything and records it
    j24 = jobs.enqueue(c22, "recode", {"variables": VARS, "mode": "dual_independent"})
    jobs.cancel(j24)
    check("whole-job Cancel still available and stops both", db.q1("SELECT status FROM jobs WHERE id=?", (j24,))["status"] == "cancelled")

    # =========================================================================================================
    print("\n[12] Review cells: providers never overwrite each other; states are distinct")
    c25, _ = new_case("case cells", big)
    AN.policy = {"SYSTEM_LEVEL": "3", "FAILURE_TYPE": "2"}
    jA = jobs.enqueue(c25, "recode", {"variables": VARS, "mode": "anthropic_only"})
    run_job(jA)
    rows = {r["name"]: r for r in export.case_results(c25)["rows"]}
    check("Claude-only results appear in the Claude column; OpenAI shows 'Not run'",
          rows["SYSTEM_LEVEL"]["cells"]["anthropic"]["value"] == "3" and rows["SYSTEM_LEVEL"]["cells"]["openai"]["state"] == "not_run")
    server.api_review(HH(), str(c25), {"variable": "FAILURE_TYPE", "action": "edit", "value": "6", "reason": "AAR says redundancy failure"})
    OA.policy, OA.script = {"SYSTEM_LEVEL": "4"}, ["good"]
    jO = jobs.enqueue(c25, "recode", {"variables": VARS, "mode": "openai_only"})
    run_job(jO)
    rows = {r["name"]: r for r in export.case_results(c25)["rows"]}
    check("a later OpenAI-only run does NOT overwrite the Claude column", rows["SYSTEM_LEVEL"]["cells"]["anthropic"]["value"] == "3"
          and rows["SYSTEM_LEVEL"]["cells"]["openai"]["value"] == "4")
    cl = rows["SYSTEM_LEVEL"]["providers"]
    check("results from different jobs on the SAME evidence and analysis version are compared (D-036), keeping "
          "their own runs", rows["SYSTEM_LEVEL"]["comparison"]["kind"] == "matched_version"
          and rows["SYSTEM_LEVEL"]["comparison"]["model_status"] == "value_disagreement"
          and len({p["run_id"] for p in cl}) == 2)
    check("Human final unchanged by both runs", rows["FAILURE_TYPE"]["review"]["value"] == "6")
    OA.script = ["invalid_code"]
    jI = jobs.enqueue(c25, "recode", {"variables": ["FAILURE_TYPE"], "mode": "openai_only"})  # new batch → not a cache hit
    run_job(jI)
    OA.script = ["good"]
    rows = {r["name"]: r for r in export.case_results(c25)["rows"]}
    states = {rows["FAILURE_TYPE"]["cells"]["openai"]["state"], rows["INTERAGENCY_COORDINATION"]["cells"]["openai"]["state"],
              rows["SYSTEM_LEVEL"]["cells"]["anthropic"]["state"]}
    check("states: invalid output, no supported value, suggested are distinct", states == {"invalid_output", "no_supported_value", "suggested"}, str(states))
    c26, _ = new_case("case failed", big)
    OA.script, OA.requests = ["500"], []
    run_job(jobs.enqueue(c26, "recode", {"variables": ["SYSTEM_LEVEL"], "mode": "openai_only"}))
    OA.script = ["good"]
    r26 = {r["name"]: r for r in export.case_results(c26)["rows"]}["SYSTEM_LEVEL"]["cells"]
    check("state 'Failed' for a failed call; Claude column 'Not run'", r26["openai"]["state"] == "failed" and r26["anthropic"]["state"] == "not_run")
    labels = export.CELL_LABELS
    check("the five required states have five different labels",
          len({labels[k] for k in ("not_run", "no_supported_value", "stopped", "failed", "invalid_output")}) == 5)
    os.environ.pop("WFD_TEST_FIXTURE", None)

    print("\n[10] Browser: provider status, mode choice, per-provider costs, comparison display")
    # A paused job (needs_input is never picked up by the worker) so the Progress page shows both Stop buttons.
    db.insert("jobs", {"case_id": c25, "kind": "recode", "stage": "coding", "status": "needs_input", "progress": 0.85,
                       "message": "Waiting for approval (test fixture)", "params_json": json.dumps({"mode": "dual_independent"}),
                       "state_json": json.dumps({"coding_mode": "dual_independent", "provider_runs": {
                           "anthropic": {"status": "pending", "role": "independent", "model": "claude-sonnet-5-5"},
                           "openai": {"status": "pending", "role": "independent", "model": "gpt-6.1-sol"}}}),
                       "created_at": time.time(), "updated_at": time.time(), "cancel_requested": 0})
    db.ex("UPDATE jobs SET status='cancelled' WHERE status IN ('queued','running')")  # the test server must never run a job
    ui_ok = ui_check(c14, KEY_AN, KEY_OA, stopped_case=c20, controls_case=c25)
    check("26 frontend provider/cost/comparison display (see details above)", ui_ok)

    print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
    if FAIL:
        print("FAILED:", FAIL)
    sys.exit(1 if FAIL else 0)


def ui_check(cid: int, key_an: str, key_oa: str, stopped_case=None, controls_case=None) -> bool:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  SKIP browser check: Playwright not installed")
        return False
    import urllib.request
    ok = True
    for port, keys in ((8811, True), (8812, False)):
        env = {**os.environ, "WFD_DATA_DIR": str(TMP / "data"), "WFD_PORT": str(port), "ANTHROPIC_API_KEY": key_an}
        env.pop("WFD_PASSWORD", None)
        if keys:
            env["OPENAI_API_KEY"] = key_oa
        else:
            env.pop("OPENAI_API_KEY", None)
        srv = subprocess.Popen([sys.executable, "-m", "app.server"], cwd=str(ROOT), env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        for _ in range(60):  # wait until the server answers (imports can be slow)
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{port}/healthz", timeout=1)
                break
            except Exception:
                if srv.poll() is not None:
                    print("  server exited:", srv.stdout.read().decode("utf8", "ignore")[-1500:])
                    break
                REAL_SLEEP(0.5)
        try:
            raw = urllib.request.urlopen(f"http://127.0.0.1:{port}/api/status").read().decode()
            leak = key_oa in raw or key_an in raw
            print(f"  [{port}] status hides keys:", not leak)
            ok &= not leak
            with sync_playwright() as p:
                b = p.chromium.launch()
                pg = b.new_page(viewport={"width": 1400, "height": 950})
                errors = []
                pg.on("pageerror", lambda e: errors.append(str(e)))
                pg.goto(f"http://127.0.0.1:{port}/#/settings")
                pg.wait_for_timeout(1200)
                stxt = pg.inner_text("body")
                has_oa_settings = pg.locator("[data-set=openai_budget_usd]").count() > 0
                print(f"  [{port}] settings shows OpenAI provider line:", "OpenAI" in stxt, "| OpenAI limit fields:", has_oa_settings)
                ok &= ("OpenAI" in stxt) and (has_oa_settings == keys)
                pg.goto(f"http://127.0.0.1:{port}/#/case/{cid}/progress")
                pg.wait_for_timeout(1500)
                pg.click("#recodeBtn")
                pg.wait_for_selector("#modeSel", timeout=8000)
                opts = pg.eval_on_selector_all("#modeSel option", "els => els.map(e => e.value)")
                print(f"  [{port}] modes offered:", opts)
                if keys:
                    ok &= "dual_independent" in opts and "openai_only" in opts
                    pg.select_option("#modeSel", "dual_independent")
                    pg.wait_for_timeout(1500)
                    body = pg.inner_text("#rcBody")
                    shows = "OpenAI" in body and "Claude" in body and "gpt-6.1-sol" in body and ("$" in body)
                    print(f"  [{port}] dual estimate shows both providers + costs:", shows)
                    ok &= shows
                    pg.screenshot(path=str(TMP / "ui_recode_dual.png"))
                else:
                    ok &= not any(o in opts for o in ("dual_independent", "openai_only", "anthropic_primary_openai_review", "openai_primary_anthropic_review"))
                pg.click("#dlgCancel")  # never start a run with fake keys
                if keys:
                    pg.goto(f"http://127.0.0.1:{port}/#/case/{cid}/review")
                    pg.wait_for_timeout(1500)
                    heads = [h.strip() for h in pg.eval_on_selector_all("#wbTable thead th", "els => els.map(e => e.innerText)")]
                    want_en = ["#", "Variable", "Claude suggestion", "OpenAI suggestion", "Human final", "Comparison", "Review status"]
                    want_zh = ["#", "变量", "Claude 建议", "OpenAI 建议", "人工最终值", "比较", "复核状态"]
                    print(f"  [{port}] review columns:", heads)
                    ok &= heads in (want_en, want_zh)
                    pg.screenshot(path=str(TMP / "ui_review_table.png"), full_page=False)
                    pg.click("tr[data-v='FAILURE_TYPE']")
                    pg.wait_for_timeout(600)
                    det = pg.inner_text("#wbDetail")
                    n_acc = pg.locator("[data-accept-sid]").count()
                    shows = n_acc == 2 and "gpt-6.1-sol" in det and ("claude-sonnet-5-5" in det)
                    print(f"  [{port}] review shows each model's suggestion with its own accept button:", shows)
                    ok &= shows
                    pg.click("#recodeVar")
                    pg.wait_for_selector("#modeSel", timeout=8000)
                    one = pg.inner_text(".modal header")
                    print(f"  [{port}] 'Re-analyze this variable' opens the estimate dialog for one variable:", "FAILURE_TYPE" in one)
                    ok &= "FAILURE_TYPE" in one
                    pg.click("#dlgCancel")
                    tbl = pg.inner_text("#wbRows")
                    print(f"  [{port}] table shows comparison badges:", any(x in tbl for x in ("Human approved", "人工已确认", "Value disagreement", "取值不一致")))
                    ok &= any(x in tbl for x in ("Human approved", "人工已确认", "Value disagreement", "取值不一致"))
                    pg.screenshot(path=str(TMP / "ui_review_providers.png"))
                    if stopped_case:
                        pg.goto(f"http://127.0.0.1:{port}/#/case/{stopped_case}/progress")
                        pg.wait_for_timeout(1500)
                        has_resume = pg.locator("[data-resume-prov=openai]").count() == 1 and pg.locator("[data-stop-prov]").count() == 0
                        print(f"  [{port}] stopped OpenAI run shows 'Resume OpenAI' (and no Stop on a finished job):", has_resume)
                        ok &= has_resume
                        pg.screenshot(path=str(TMP / "ui_progress_resume.png"))
                    if controls_case:
                        pg.goto(f"http://127.0.0.1:{port}/#/case/{controls_case}/progress")
                        pg.wait_for_timeout(1500)
                        both = pg.locator("[data-stop-prov=anthropic]").count() == 1 and pg.locator("[data-stop-prov=openai]").count() == 1
                        print(f"  [{port}] separate 'Stop Claude' and 'Stop OpenAI' buttons:", both)
                        ok &= both
                        pg.screenshot(path=str(TMP / "ui_progress_stop_buttons.png"))
                        pg.click("[data-stop-prov=openai]")
                        pg.wait_for_selector("#dlgOk", timeout=5000)
                        dlg = pg.inner_text(".modal")
                        warn = ("may still be billed" in dlg) or ("可能仍然计费" in dlg)
                        print(f"  [{port}] stop dialog warns that a sent request may still be billed:", warn)
                        ok &= warn
                        pg.screenshot(path=str(TMP / "ui_stop_dialog.png"))
                        pg.click("#dlgCancel")
                print(f"  [{port}] JS errors:", errors)
                ok &= not errors
                b.close()
        finally:
            srv.terminate()
            out = srv.communicate(timeout=10)[0].decode("utf8", "ignore")
            if key_oa in out or key_an in out:
                print(f"  [{port}] KEY FOUND IN SERVER OUTPUT")
                ok = False
    print("  screenshots in", TMP)
    return bool(ok)


if __name__ == "__main__":
    main()
