"""Missing values: blank vs -9 vs N/A (D-039). Run:  python -m tests.test_missingness

Offline (no network, no model). Checks every missingness behaviour that is implemented from the codebook:
-9 only where a variable's own entry defines it (for every variable, every type), placeholders become blank, the
codebook's -9 wording reaches the prompt, workbook -9/placeholders are flagged (never adopted), codebook consistency
rules on human-final values, INCIDENT_DURATION never derived from reversed or missing dates, and the audit table.
"""
from __future__ import annotations

import csv
import io
import json
import os
import socket
import sys
import tempfile
import time
from pathlib import Path

TMP = Path(tempfile.mkdtemp(prefix="wfdmiss_"))
os.environ["WFD_DATA_DIR"] = str(TMP / "data")
for k in ("ANTHROPIC_API_KEY", "TAVILY_API_KEY", "BRAVE_API_KEY", "OPENAI_API_KEY", "OPENAI_BASE_URL", "WFD_TEST_FIXTURE"):
    os.environ.pop(k, None)
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
PASS, FAIL, NET = [], [], []
_real = socket.socket.connect


def _guard(self, addr):
    host = addr[0] if isinstance(addr, tuple) else str(addr)
    if host not in ("127.0.0.1", "localhost", "::1"):
        NET.append(host)
        raise OSError("network blocked in tests")
    return _real(self, addr)


socket.socket.connect = _guard


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print(("  PASS " if cond else "  FAIL ") + name + (f"  — {detail}" if detail and not cond else ""))


Q = "The county said the delay could not be determined from the system logs"
TEXT = f"Kerr County, Texas, July 4, 2025. {Q}. Officials gave no figure for message length. "
DEFINED = {"TRANSMISSION_LATENCY", "DELIVERY_COVERAGE", "POPULATION_AFFECTED_ESTIMATE", "HARM_LOSS_INDICATORS",
           "MESSAGE_DISSEMINATION_TIME", "SOURCE_COUNT_PRIMARY", "SOURCE_COUNT_SECONDARY"}


def main():
    from app import coding, db, export, research, server
    from app.missingness import consistency_flags, duration_days
    from app.validator import validate
    from tools import missingness_audit as audit
    db.conn()
    server.bootstrap_schema()
    schema = coding.active_schema()
    F = {f["name"]: f for f in schema["fields"]}
    P = {"S1-P1": {"text": TEXT, "source_id": 1}}
    ev = [{"id": "S1-P1", "quote": Q, "stance": "supports"}]

    def v(name, value, evidence=ev, options=None):
        return validate(F[name], {"value": value, "evidence": evidence, "options": options or []}, P)

    print("\n[1] -9 only where the variable's own codebook entry defines it")
    check("1 exactly the 7 codebook variables define -9",
          {n for n, f in F.items() if "-9" in f.get("missing_codes", [])} == DEFINED)
    for n in sorted(DEFINED):
        r = v(n, "-9")
        check(f"2 {n}: -9 with a cited passage → valid, with a 'confirm' warning",
              r["ok"] and r["value"] == "-9" and any("special missing value" in w for w in r["warnings"]), str(r["errors"]))
    r = v("TRANSMISSION_LATENCY", "-9.0")
    check("3 Excel-style '-9.0' is the same special value (stored as '-9')", r["ok"] and r["value"] == "-9")
    r = v("TRANSMISSION_LATENCY", "-9", evidence=[])
    check("3b -9 still needs a cited supporting passage", not r["ok"] and "no_supporting_evidence" in r["reason_codes"])
    for n, label in (("MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)", "numeric (was accepted as the number -9)"),
                     ("REGULATORY_OR_POLICY_FRAMEWORK", "free text"), ("ALERT_ORIGINATOR_PLATFORM", "open list"),
                     ("END_DATE", "date"), ("GEOCODE_SPECIFICITY", "single-select code"),
                     ("INCIDENT_DURATION", "derived numeric")):
        r = v(n, "-9")
        check(f"4 {n} ({label}): -9 rejected with reason missing_code_not_defined",
              not r["ok"] and r["value"] == "" and r["reason_codes"] == ["missing_code_not_defined"], str(r)[:200])
    r = v("MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)", "-9.0")
    check("4b '-9.0' also rejected where -9 is not defined", r["reason_codes"] == ["missing_code_not_defined"])
    r = v("SYSTEM_INVOLVED", "WEA, -9", evidence=[], options=[{"code": "WEA", "evidence": ev}, {"code": "-9", "evidence": ev}])
    sel = {s["code"]: s for s in r["selections"]}
    check("5 multi-select: -9 selection rejected (missing_code_not_defined), the supported code kept for review",
          r["outcome"] == "partially_valid" and r["value"] == "WEA" and "missing_code_not_defined" in sel["-9"]["reasons"])
    r = v("TRANSMISSION_LATENCY", "5")
    check("5b ordinary numbers are unaffected", r["ok"] and r["value"] == "5")
    r = validate(F["TRANSMISSION_LATENCY"], {"value": "-9", "evidence": ev}, P)
    r2 = validate(F["POPULATION_AFFECTED_ESTIMATE"], {"value": "-9, 500", "evidence": ev}, P)
    check("5c -9 cannot be mixed with a value", r["ok"] and not r2["ok"])
    model_coded = [f for f in schema["fields"] if not f.get("rule_missing") and f["field_class"] in ("sourced", "judgment")]
    wrong = [f["name"] for f in model_coded if v(f["name"], "-9")["ok"] != (f["name"] in DEFINED)]
    check(f"6 every model-coded variable ({len(model_coded)}): -9 accepted iff the codebook defines it", not wrong, str(wrong))

    print("\n[2] Placeholders and the codebook's 'leave blank' rules")
    for n in ("LAT/LONG", "CITY_OR_TOWN", "COUNTY_OR_PARISH"):
        for ph in ("-", "N/A", "—"):
            r = v(n, ph)
            if not (r["ok"] and r["value"] == "" and r["warnings"]):
                check(f"7 {n}: placeholder {ph!r} → blank with a warning", False, str(r)[:200])
                break
        else:
            check(f"7 {n}: placeholders '-', 'N/A', '—' → blank with a warning", True)
    r = v("LAT/LONG", "Latitude: 21.30° N, Longitude: -157.85° W")
    check("7b a real negative coordinate is kept", r["ok"] and r["value"].endswith("-157.85° W"))

    print("\n[3] Codebook wording reaches the prompt; workbook -9 is flagged, never adopted")
    check("8 TRANSMISSION_LATENCY -9 label is the codebook wording 'unknown or inapplicable'",
          F["TRANSMISSION_LATENCY"]["missing_labels"].get("-9") == "unknown or inapplicable")
    check("8b MESSAGE_DISSEMINATION_TIME -9 label 'Unknown or not documented'",
          F["MESSAGE_DISSEMINATION_TIME"]["missing_labels"].get("-9") == "Unknown or not documented")
    spec = coding.field_spec_text(F["TRANSMISSION_LATENCY"])
    spec2 = coding.field_spec_text(F["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"])
    check("9 prompt: -9 shown with codebook wording where defined; 'leave blank' where not",
          "-9 = unknown or inapplicable" in spec and "No special missing value is defined" in spec2)
    iss = " ".join(F["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"]["issues"])
    check("10 workbook -9 in a numeric field without -9 is flagged (35 rows)", "uses -9 in 35 row(s)" in iss, iss[:200])
    check("10b workbook -9 in a categorical field is flagged", any("'-9'" in i for i in F["ALERT_APPROVAL_PROCESS"]["issues"]))
    check("10c workbook placeholders flagged (LAT/LONG '-')", any("placeholder" in i for i in F["LAT/LONG"]["issues"]))
    check("10d workbook flags never expand the allowed values",
          F["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"]["missing_codes"] == []
          and "-9" not in {c["code"] for c in F["ALERT_APPROVAL_PROCESS"]["codes"]})
    check("10e workbook-data notes are not sent to the model as rules",
          "Workbook data" not in coding.field_spec_text(F["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"]))

    print("\n[4] Codebook consistency rules on human-final values (flags only)")
    fl = consistency_flags({"UNCERTAINTY_FLAG": "1", "UNCERTAINTY_TYPE": ""})
    check("11 UNCERTAINTY_FLAG = 1 with blank UNCERTAINTY_TYPE → flagged", "UNCERTAINTY_TYPE" in fl)
    check("11b UNCERTAINTY_FLAG = 0 with blank type → no flag (blank for flag 0 is PI-M7, not enforced)",
          not consistency_flags({"UNCERTAINTY_FLAG": "0", "UNCERTAINTY_TYPE": ""}))
    check("11c Excel-style '1.0' read as 1", "UNCERTAINTY_TYPE" in consistency_flags({"UNCERTAINTY_FLAG": "1.0",
                                                                                      "UNCERTAINTY_TYPE": ""}))
    check("12 blank LAT/LONG with GEOCODE_SPECIFICITY 3 → flagged; with 0 → not flagged",
          "GEOCODE_SPECIFICITY" in consistency_flags({"LAT/LONG": "", "GEOCODE_SPECIFICITY": "3"})
          and not consistency_flags({"LAT/LONG": "", "GEOCODE_SPECIFICITY": "0"}))
    fl = consistency_flags({"EVENT_DATE": "2025-07-04", "END_DATE": "2025-07-01"})
    check("13 END_DATE before EVENT_DATE → both flagged", set(fl) == {"EVENT_DATE", "END_DATE"})
    fl = consistency_flags({"EVENT_DATE": "2022-06-10", "END_DATE": "2022-06-12", "INCIDENT_DURATION": "2"})
    check("14 INCIDENT_DURATION differing from (END – EVENT) + 1 → flagged with the codebook result",
          "INCIDENT_DURATION" in fl and "gives 3" in fl["INCIDENT_DURATION"][0])
    check("14b consistent record → no flags", not consistency_flags({"EVENT_DATE": "2022-06-10", "END_DATE": "2022-06-12",
                                                                     "INCIDENT_DURATION": "3", "UNCERTAINTY_FLAG": "0"}))
    check("15 duration_days follows the codebook example (2022-06-10 → 2022-06-12 = 3) and refuses reversed/missing",
          duration_days("2022-06-10", "2022-06-12") == 3 and duration_days("2022-06-12", "2022-06-10") is None
          and duration_days("2022-06-10", "") is None)

    print("\n[5] Through the case: derivation, review flags, export")
    settings = {**coding.DEFAULT_SETTINGS} if hasattr(coding, "DEFAULT_SETTINGS") else {}
    cid = db.insert("cases", {"name": "Test flood", "location": "Kerr County, Texas", "date_text": "July 4, 2025", "details": "",
                              "known_links": "", "aliases": "", "schema_version_id": schema["id"], "status": "new",
                              "identity_json": "{}", "settings_json": json.dumps(settings), "created_at": time.time(),
                              "updated_at": time.time()})
    research.ingest_manual(cid, "text", {"title": "County statement", "text": TEXT * 3})
    H = type("H", (), {"user": "coder1"})()

    def review(var, value):
        server.api_review(H, str(cid), {"variable": var, "action": "edit", "value": value, "reason": "test"})
    case = db.q1("SELECT * FROM cases WHERE id=?", (cid,))
    target = [F["INCIDENT_DURATION"]]

    def derived():
        coding.derive_fields(case, schema, None, target)
        return db.q1("SELECT value, rationale FROM suggestions WHERE case_id=? AND variable='INCIDENT_DURATION' "
                     "ORDER BY id DESC", (cid,))
    review("EVENT_DATE", "2025-07-04")
    d = derived()
    check("16 END_DATE missing → INCIDENT_DURATION blank (never -9)", d["value"] == "", str(d))
    review("END_DATE", "2025-07-01")
    d = derived()
    check("17 END_DATE before EVENT_DATE → INCIDENT_DURATION blank with an explanation",
          d["value"] == "" and "earlier than" in d["rationale"], str(d))
    review("END_DATE", "2025-07-06")
    d = derived()
    check("18 valid dates → INCIDENT_DURATION = (END – EVENT) + 1 = 3", d["value"] == "3", str(d))
    review("UNCERTAINTY_FLAG", "1")
    review("LAT/LONG", "")
    review("GEOCODE_SPECIFICITY", "3")
    rows = {r["name"]: r for r in export.case_results(cid)["rows"]}
    check("19 review workbench rows carry the rule flags (UNCERTAINTY_TYPE not yet reviewed → no flag yet)",
          rows["GEOCODE_SPECIFICITY"]["rule_flags"] and not rows["UNCERTAINTY_TYPE"]["rule_flags"])
    review("UNCERTAINTY_TYPE", "")
    rows = {r["name"]: r for r in export.case_results(cid)["rows"]}
    check("19b human-final blank UNCERTAINTY_TYPE with flag 1 → flagged", bool(rows["UNCERTAINTY_TYPE"]["rule_flags"]))
    review("MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)", "-9")
    rows = {r["name"]: r for r in export.case_results(cid)["rows"]}
    check("19c a researcher's -9 where the codebook defines none is kept but flagged (PI-M3), not blocked",
          any("PI-M3" in x for x in rows["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"]["rule_flags"])
          and rows["MESSAGE_CHARACTERISTICS_LENGTH (Char. No Space)"]["review"]["value"] == "-9")
    review("TRANSMISSION_LATENCY", "-9")
    check("19d a researcher's -9 where the codebook defines it → no flag",
          not {r["name"]: r for r in export.case_results(cid)["rows"]}["TRANSMISSION_LATENCY"]["rule_flags"])
    check("20 flags never change human-final values",
          db.q1("SELECT value FROM reviews WHERE case_id=? AND variable='GEOCODE_SPECIFICITY'", (cid,))["value"] == "3")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(export.xlsx(cid)))
    hdr = [c.value for c in wb["Results"][1]]
    res = {r[0]: dict(zip(hdr, r)) for r in wb["Results"].iter_rows(min_row=2, values_only=True)}
    check("21 export: 'Rule checks' column and explanation carry the flag; the value is exported unchanged",
          "Rule checks" in hdr and "GEOCODE_SPECIFICITY = 0" in (res["GEOCODE_SPECIFICITY"]["Rule checks"] or "")
          and str(res["GEOCODE_SPECIFICITY"]["Value"]) == "3" and "rule check" in res["GEOCODE_SPECIFICITY"]["Explanation"])
    check("21b workbook-aligned Case_Row keeps one column per workbook header",
          len([c.value for c in wb["Case_Row"][1]]) == len(schema["fields"]))

    print("\n[6] Audit table (machine-readable)")
    rows_a = audit.main(write=False)
    names = {r["variable"] for r in rows_a}
    check("22 audit covers every workbook column (82) with the requested columns",
          len(rows_a) == len(schema["fields"]) == 82 and all(set(audit.COLUMNS) == set(r) for r in rows_a))
    focus = ["TRANSMISSION_LATENCY", "DELIVERY_COVERAGE", "POPULATION_AFFECTED_ESTIMATE", "HARM_LOSS_INDICATORS",
             "MESSAGE_DISSEMINATION_TIME", "END_DATE", "INCIDENT_DURATION"]
    A = {r["variable"]: r for r in rows_a}
    check("23 benchmark-disagreement variables are present, documented and sent for a PI decision",
          all(n in names and A[n]["pi_decision_required"] == "yes" and A[n]["conflict_ambiguity"] != "None found."
              for n in focus))
    check("24 audit '-9 allowed?' matches the codebook-derived schema for every variable",
          all((A[n]["minus9_allowed"] == "yes") == ("-9" in F[n].get("missing_codes", [])) for n in F))
    flagged = [r["variable"] for r in rows_a if r["minus9_allowed"] == "no" and r["workbook_minus9_rows"]]
    check("25 every workbook -9 where the codebook defines none is flagged for a PI decision (21 variables)",
          len(flagged) == 21 and all("PI-M3" in A[n]["pi_decision_ids"] for n in flagged), str(flagged))
    with open(ROOT / "docs" / "missingness" / "missingness_audit.csv") as fh:
        saved = list(csv.DictReader(fh))
    fresh = [{k: str(v) for k, v in r.items()} for r in rows_a]
    check("26 the committed audit CSV matches a fresh regeneration (no stale table)", saved == fresh)
    pi_doc = (ROOT / "docs" / "missingness" / "PI_DECISIONS_missingness.md").read_text()
    ids = {i.strip() for r in rows_a for i in r["pi_decision_ids"].split(",") if i.strip()}
    check("27 every PI decision id used in the table is explained in the PI decision document",
          all(f"## {i} " in pi_doc for i in ids), str(sorted(i for i in ids if f"## {i} " not in pi_doc)))
    check("28 no network access was attempted", not NET, str(NET))
    print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
    if FAIL:
        print("FAILED:", *FAIL, sep="\n  ")
        sys.exit(1)


if __name__ == "__main__":
    main()
